"""
AI Layer feature for generating Description and Recommendation fields in Excel.

Users upload:
- An Excel workbook with target narrative columns empty
- A ZIP file containing POC screenshots

The service iterates each vulnerability row, matches screenshots, calls OpenAI,
and writes the generated narrative fields back into the workbook.
"""

from __future__ import annotations

import base64
import json
import os
import re
import shutil
import tempfile
import time
import traceback
import zipfile
from pathlib import Path
from typing import Any

# Load .env so the sub-app always picks up OPENAI_API_KEY regardless of import order
from dotenv import load_dotenv as _load_dotenv
_load_dotenv(Path(__file__).resolve().parent / ".env", override=True)

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse
from openai import OpenAI
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from PIL import Image
from sqlalchemy.orm import Session
from starlette.background import BackgroundTask

from db import AuditLog, SessionLocal
from .ai_prompts import SYSTEM_PROMPT, build_row_payload, build_user_prompt, parse_ai_json, validate_ai_output

app = FastAPI()

VALID_REPORT_TYPES = {"ip", "web", "certin", "mitkat"}
MAX_IMAGES_PER_ROW = 4
MAX_IMAGE_WIDTH = 1024
REQUEST_DELAY_SECONDS = 0.4
RATE_LIMIT_WINDOW_SECONDS = 600
RATE_LIMIT_MAX_REQUESTS = 5

CORE_HEADERS = {
    "Observation/Vulnerability Title": "vulnerability_title",
    "Severity": "severity",
    "Detailed Observation/ Vulnerable Point": "detailed_observation",
    "Recommendation": "recommendation",
}

OPTIONAL_HEADERS = {
    "Asset/Hostname": "asset_hostname",
    "Tester_Name": "tester_name",
    "Project": "project",
    "Sr no.": "row_id",
    "Observation": "observation_number",
    "Status": "status",
    "New or Repeat Observation": "new_or_repeat",
    "CVE/CWE": "cve_cwe",
    "CVSS": "cvss",
    "CVSS Version Ref": "cvss_vector",
    "Affected Asset i.e. IP/URL/Application etc.": "affected_asset",
    "Reference": "reference",
    "Evidence / Proof of Concept": "evidence",
    "Screenshot1": "screenshot1",
    "Screenshot2": "screenshot2",
    "Screenshot3": "screenshot3",
    "Screenshot4": "screenshot4",
}

ALL_KNOWN_HEADERS = {**CORE_HEADERS, **OPTIONAL_HEADERS}

OUTPUT_HEADERS = {
    "detailed_observation": "Detailed Observation/ Vulnerable Point",
    "recommendation": "Recommendation",
    "reference": "Reference",
    "evidence": "Evidence / Proof of Concept",
}

TYPE_LABELS = {
    "ip": "IP_Report",
    "web": "Web_Report",
    "certin": "CertIN_Report",
    "mitkat": "MitKat_Report",
}

_REQUEST_TIMESTAMPS_BY_USER: dict[str, list[float]] = {}

OUTPUT_HEADER_KEYWORDS = {
    "detailed_observation": ("detailedobservation", "vulnerablepoint"),
    "recommendation": ("recommendation",),
    "reference": ("reference",),
    "evidence": ("evidence", "proofofconcept"),
}


class OpenAIQuotaError(Exception):
    """OpenAI billing/quota exhausted (insufficient_quota). Abort early with a clear message."""

    pass


def _is_insufficient_quota(exc: BaseException) -> bool:
    text = str(exc).lower()
    if "insufficient_quota" in text:
        return True
    body = getattr(exc, "body", None)
    if isinstance(body, dict):
        err = body.get("error") or {}
        if err.get("code") == "insufficient_quota":
            return True
    return False


def _read_html_template() -> str:
    html_path = Path(__file__).resolve().parent.parent / "ai_layer.html"
    with open(html_path, encoding="utf-8") as handle:
        return handle.read()


def _normalize_header(value: Any) -> str:
    return str(value or "").strip()


def _normalize_key(value: Any) -> str:
    text = str(value or "").strip().lower()
    return re.sub(r"[^a-z0-9#._-]+", "", text)


def _normalize_path_part(value: str) -> str:
    text = value.strip().lower()
    text = text.replace("\\", "/").split("/")[-1]
    return re.sub(r"[^a-z0-9#._-]+", "", text)


def _extract_host_candidate(value: str) -> str:
    cleaned = value.strip()
    if "://" in cleaned:
        cleaned = cleaned.split("://", 1)[1]
    cleaned = cleaned.split("/", 1)[0].split(":", 1)[0]
    return _normalize_path_part(cleaned)


def _get_openai_client() -> OpenAI:
    api_key = (os.getenv("OPENAI_API_KEY") or "").strip()
    if not api_key:
        raise HTTPException(status_code=500, detail="OPENAI_API_KEY is not configured")
    kwargs: dict[str, str] = {"api_key": api_key}
    org = (os.getenv("OPENAI_ORGANIZATION") or os.getenv("OPENAI_ORG_ID") or "").strip()
    if org:
        kwargs["organization"] = org
    project = (os.getenv("OPENAI_PROJECT") or os.getenv("OPENAI_PROJECT_ID") or "").strip()
    if project:
        kwargs["project"] = project
    return OpenAI(**kwargs)


def _build_header_map(worksheet) -> dict[str, int]:
    header_map: dict[str, int] = {}
    for column_index, cell in enumerate(worksheet[1], start=1):
        header_name = _normalize_header(cell.value)
        if header_name:
            header_map[header_name] = column_index

    missing_core = [header for header in CORE_HEADERS if header not in header_map]
    if missing_core:
        raise HTTPException(
            status_code=400,
            detail=f"Uploaded Excel is missing required columns: {', '.join(missing_core)}",
        )
    return header_map


def _cell_value(worksheet, row_index: int, header_map: dict[str, int], header: str) -> str:
    col_idx = header_map.get(header)
    if col_idx is None:
        return ""
    return str(worksheet.cell(row=row_index, column=col_idx).value or "").strip()


def _row_to_payload(
    worksheet,
    row_index: int,
    header_map: dict[str, int],
    report_type: str,
) -> dict[str, str]:
    payload: dict[str, str] = {"report_type": report_type}
    for header, payload_key in ALL_KNOWN_HEADERS.items():
        payload[payload_key] = _cell_value(worksheet, row_index, header_map, header)
    return payload


def _should_generate_row(row_payload: dict[str, str]) -> bool:
    if not row_payload.get("vulnerability_title"):
        return False

    target_fields = (
        row_payload.get("detailed_observation", ""),
        row_payload.get("recommendation", ""),
        row_payload.get("reference", ""),
        row_payload.get("evidence", ""),
    )
    return any(not value.strip() for value in target_fields)


def _extract_zip(zip_path: str, extract_dir: str) -> list[Path]:
    with zipfile.ZipFile(zip_path, "r") as archive:
        archive.extractall(extract_dir)

    image_paths = [
        path for path in Path(extract_dir).rglob("*")
        if path.is_file() and path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}
    ]
    return sorted(image_paths)


def _score_image_path(path: Path, row_payload: dict[str, str], screenshot_names: list[str]) -> int:
    parts = [_normalize_path_part(part) for part in path.parts]
    filename = _normalize_path_part(path.name)
    score = 0

    row_id = _normalize_key(row_payload.get("row_id"))
    observation_number = _normalize_key(row_payload.get("observation_number"))
    severity = _normalize_key(row_payload.get("severity"))
    asset_hostname = _extract_host_candidate(row_payload.get("asset_hostname", ""))
    affected_asset = _extract_host_candidate(row_payload.get("affected_asset", ""))

    if row_id and row_id in parts:
        score += 10
    if observation_number and observation_number in parts:
        score += 10
    if severity and severity in parts:
        score += 3
    if asset_hostname and asset_hostname in parts:
        score += 4
    if affected_asset and affected_asset in parts:
        score += 4
    if screenshot_names and filename in screenshot_names:
        score += 8

    return score


def _find_matching_screenshots(image_paths: list[Path], row_payload: dict[str, str]) -> list[Path]:
    screenshot_names = [
        _normalize_path_part(row_payload.get(key, ""))
        for key in ("screenshot1", "screenshot2", "screenshot3", "screenshot4")
        if row_payload.get(key)
    ]

    scored_paths: list[tuple[int, Path]] = []
    for image_path in image_paths:
        score = _score_image_path(image_path, row_payload, screenshot_names)
        if score > 0:
            scored_paths.append((score, image_path))

    if not scored_paths:
        return []

    scored_paths.sort(key=lambda item: (-item[0], str(item[1]).lower()))
    selected = [path for _, path in scored_paths[:MAX_IMAGES_PER_ROW]]
    selected.sort(key=lambda path: str(path).lower())
    return selected


def _prepare_image_part(image_path: Path) -> dict[str, Any]:
    with Image.open(image_path) as image:
        image = image.convert("RGB")
        if image.width > MAX_IMAGE_WIDTH:
            ratio = MAX_IMAGE_WIDTH / float(image.width)
            new_size = (MAX_IMAGE_WIDTH, max(1, int(image.height * ratio)))
            image = image.resize(new_size)

        temp_image = tempfile.NamedTemporaryFile(delete=False, suffix=".jpg")
        temp_image_path = temp_image.name
        temp_image.close()
        image.save(temp_image_path, format="JPEG", quality=85)

    try:
        with open(temp_image_path, "rb") as handle:
            encoded = base64.b64encode(handle.read()).decode("utf-8")
    finally:
        if os.path.exists(temp_image_path):
            os.remove(temp_image_path)

    return {
        "type": "image_url",
        "image_url": {"url": f"data:image/jpeg;base64,{encoded}", "detail": "low"},
    }


def _build_message_content(row_payload: dict[str, str], image_paths: list[Path]) -> list[dict[str, Any]]:
    screenshot_names = [path.name for path in image_paths]
    content: list[dict[str, Any]] = [
        {"type": "text", "text": build_user_prompt(build_row_payload(row_payload), screenshot_names)}
    ]
    for image_path in image_paths:
        content.append(_prepare_image_part(image_path))
    return content


def _call_openai_for_row(client: OpenAI, row_payload: dict[str, str], image_paths: list[Path]) -> dict[str, str]:
    model = os.getenv("OPENAI_MODEL", "gpt-4o").strip() or "gpt-4o"
    response = client.chat.completions.create(
        model=model,
        temperature=0.2,
        response_format={"type": "json_object"},
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": _build_message_content(row_payload, image_paths)},
        ],
    )
    content = response.choices[0].message.content or ""
    parsed = parse_ai_json(content)
    return validate_ai_output(parsed, row_payload, len(image_paths))


def _generate_row_with_retries(
    client: OpenAI,
    row_payload: dict[str, str],
    image_paths: list[Path],
    retries: int = 2,
) -> dict[str, str]:
    last_error: Exception | None = None
    for attempt in range(retries + 1):
        try:
            return _call_openai_for_row(client, row_payload, image_paths)
        except Exception as exc:  # pragma: no cover - defensive retry boundary
            if _is_insufficient_quota(exc):
                raise OpenAIQuotaError(
                    "OpenAI returned insufficient_quota. Add prepaid credits, wait for billing to activate, "
                    "revoke exposed keys, create a NEW secret key in the same org as your balance, "
                    "and see docs/OPENAI_AI_LAYER_PRODUCTION.md"
                ) from exc
            last_error = exc
            if attempt < retries:
                time.sleep(1.0)
    raise RuntimeError(str(last_error) if last_error else "Unknown AI generation error")


def _write_ai_output(
    worksheet,
    row_index: int,
    output_col_map: dict[str, int],
    generated: dict[str, str],
) -> None:
    for output_key in OUTPUT_HEADERS:
        col_idx = output_col_map.get(output_key)
        if col_idx is not None:
            worksheet.cell(row=row_index, column=col_idx).value = generated[output_key]


def _cleanup_temp_dir(temp_dir: str) -> None:
    if os.path.exists(temp_dir):
        shutil.rmtree(temp_dir, ignore_errors=True)


def _enforce_rate_limit(request: Request) -> None:
    user = request.session.get("user") or {}
    user_key = (user.get("email") or request.client.host if request.client else "anonymous") or "anonymous"
    now = time.time()
    window_start = now - RATE_LIMIT_WINDOW_SECONDS
    history = _REQUEST_TIMESTAMPS_BY_USER.setdefault(user_key, [])
    history[:] = [ts for ts in history if ts >= window_start]
    if len(history) >= RATE_LIMIT_MAX_REQUESTS:
        raise HTTPException(
            status_code=429,
            detail=(
                f"Rate limit exceeded for AI generation. Max {RATE_LIMIT_MAX_REQUESTS} requests per "
                f"{RATE_LIMIT_WINDOW_SECONDS // 60} minutes. Please wait and retry."
            ),
        )
    history.append(now)


def _resolve_output_columns(worksheet, header_map: dict[str, int]) -> dict[str, int]:
    resolved: dict[str, int] = {}

    # Exact matches first.
    for key, header in OUTPUT_HEADERS.items():
        if header in header_map:
            resolved[key] = header_map[header]

    # Fallback: keyword/partial matching on normalized headers.
    if len(resolved) < len(OUTPUT_HEADERS):
        for col_idx, cell in enumerate(worksheet[1], start=1):
            normalized = re.sub(r"[^a-z0-9]+", "", str(cell.value or "").lower())
            if not normalized:
                continue
            for key, keywords in OUTPUT_HEADER_KEYWORDS.items():
                if key in resolved:
                    continue
                if all(k in normalized for k in keywords):
                    resolved[key] = col_idx
                    break

    return resolved


def _autofit_worksheet(worksheet, output_col_map: dict[str, int]) -> None:
    if not output_col_map:
        return

    desired_widths = {
        "detailed_observation": 55,
        "recommendation": 55,
        "reference": 45,
        "evidence": 55,
    }
    col_widths: dict[int, float] = {}

    for key, col_idx in output_col_map.items():
        col_letter = get_column_letter(col_idx)
        width = desired_widths.get(key, 50)
        worksheet.column_dimensions[col_letter].width = width
        col_widths[col_idx] = width

        for row_idx in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top",
                horizontal=cell.alignment.horizontal if cell.alignment else "left",
            )

    # Compute row heights so full wrapped content is visible.
    for row_idx in range(2, worksheet.max_row + 1):
        required_lines = 1
        for col_idx, width in col_widths.items():
            text = str(worksheet.cell(row=row_idx, column=col_idx).value or "").strip()
            if not text:
                continue
            chars_per_line = max(16, int(width * 1.05))
            logical_lines = 0
            for segment in (text.splitlines() or [text]):
                logical_lines += max(1, (len(segment) + chars_per_line - 1) // chars_per_line)
            required_lines = max(required_lines, logical_lines)

        # No tight cap: keep content visible without manual expand.
        worksheet.row_dimensions[row_idx].height = max(24, min(1200, required_lines * 16))


def _log_audit(request: Request, action: str, metadata: dict[str, Any]) -> None:
    db: Session = SessionLocal()
    try:
        user = request.session.get("user") or {}
        db.add(
            AuditLog(
                user_email=user.get("email"),
                user_name=user.get("name"),
                action=action,
                metadata_json=json.dumps(metadata),
                ip_address=request.client.host if request.client else None,
                user_agent=request.headers.get("user-agent"),
            )
        )
        db.commit()
    except Exception:
        db.rollback()
    finally:
        db.close()


@app.get("/", response_class=HTMLResponse)
async def ai_layer_home() -> HTMLResponse:
    return HTMLResponse(_read_html_template())


@app.post("/generate-excel/")
async def generate_excel_with_ai(
    request: Request,
    excel_file: UploadFile = File(...),
    poc_zip: UploadFile = File(...),
    report_type: str = Form(...),
):
    report_type = report_type.strip().lower()
    if report_type not in VALID_REPORT_TYPES:
        raise HTTPException(status_code=400, detail="Invalid report type")

    if not excel_file.filename or not excel_file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Please upload a valid .xlsx file")
    if not poc_zip.filename or not poc_zip.filename.lower().endswith(".zip"):
        raise HTTPException(status_code=400, detail="Please upload a valid .zip POC file")
    _enforce_rate_limit(request)

    temp_dir = tempfile.mkdtemp(prefix="ai_layer_")
    excel_path = os.path.join(temp_dir, excel_file.filename)
    zip_path = os.path.join(temp_dir, poc_zip.filename)
    extract_dir = os.path.join(temp_dir, "poc")

    try:
        with open(excel_path, "wb") as handle:
            handle.write(await excel_file.read())
        with open(zip_path, "wb") as handle:
            handle.write(await poc_zip.read())

        workbook = load_workbook(excel_path)
        worksheet = workbook[workbook.sheetnames[0]]
        header_map = _build_header_map(worksheet)
        output_col_map = _resolve_output_columns(worksheet, header_map)
        image_paths = _extract_zip(zip_path, extract_dir)
        client = _get_openai_client()

        processed_rows = 0
        skipped_rows = 0
        failed_rows: list[str] = []

        for row_index in range(2, worksheet.max_row + 1):
            row_payload = _row_to_payload(worksheet, row_index, header_map, report_type)
            if not _should_generate_row(row_payload):
                skipped_rows += 1
                continue

            matched_images = _find_matching_screenshots(image_paths, row_payload)
            try:
                generated = _generate_row_with_retries(client, row_payload, matched_images)
                _write_ai_output(worksheet, row_index, output_col_map, generated)
                processed_rows += 1
            except OpenAIQuotaError:
                _cleanup_temp_dir(temp_dir)
                raise HTTPException(
                    status_code=503,
                    detail=(
                        "OpenAI insufficient_quota: your API key cannot spend credits yet. "
                        "Create a new secret key after credits are active, confirm the key's org matches "
                        "Home credit balance, and read docs/OPENAI_AI_LAYER_PRODUCTION.md"
                    ),
                )
            except Exception as row_exc:
                failed_rows.append(row_payload.get("row_id") or f"row-{row_index}")
                fail_msg = "AI generation failed for this row. Please review manually."
                for out_key in ("detailed_observation", "recommendation", "evidence"):
                    col_idx = output_col_map.get(out_key)
                    if col_idx is not None:
                        worksheet.cell(row=row_index, column=col_idx).value = fail_msg
                ref_col = output_col_map.get("reference")
                if ref_col is not None:
                    worksheet.cell(row=row_index, column=ref_col).value = ""
                print(f"AI row generation failed for {failed_rows[-1]}: {row_exc}")

            if REQUEST_DELAY_SECONDS > 0:
                time.sleep(REQUEST_DELAY_SECONDS)

        if processed_rows == 0:
            if failed_rows:
                raise HTTPException(
                    status_code=502,
                    detail=f"AI generation failed for all {len(failed_rows)} eligible row(s): {', '.join(failed_rows[:10])}. Check your OPENAI_API_KEY, billing, and model access.",
                )
            raise HTTPException(
                status_code=400,
                detail="No rows were eligible for AI generation. Make sure the target columns are blank and vulnerability titles are present.",
            )

        _autofit_worksheet(worksheet, output_col_map)

        type_label = TYPE_LABELS.get(report_type, report_type)
        output_filename = f"{type_label}_ai_completed.xlsx"
        output_path = os.path.join(temp_dir, output_filename)
        workbook.save(output_path)

        _log_audit(
            request,
            "ai-layer-generate-excel",
            {
                "report_type": report_type,
                "processed_rows": processed_rows,
                "skipped_rows": skipped_rows,
                "failed_rows": failed_rows,
                "images_found": len(image_paths),
                "output_filename": output_filename,
            },
        )

        return FileResponse(
            path=output_path,
            filename=output_filename,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            background=BackgroundTask(_cleanup_temp_dir, temp_dir),
        )
    except HTTPException:
        _cleanup_temp_dir(temp_dir)
        raise
    except Exception as exc:
        traceback.print_exc()
        _cleanup_temp_dir(temp_dir)
        raise HTTPException(status_code=500, detail=f"AI Excel generation failed: {exc}") from exc


@app.get("/health")
async def ai_layer_health() -> JSONResponse:
    return JSONResponse({"status": "ok", "service": "ai-layer"})
