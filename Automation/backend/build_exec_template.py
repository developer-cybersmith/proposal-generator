"""
One-time / CI utility: build exec_summary_template.docx with CyberSmithSECURE branding.

Run from repo root:
  python Automation/backend/build_exec_template.py

Requires assets/cover_bg.png, header_logo.png, footer_banner.png (extracted from CSS reference).
"""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BACKEND_DIR = Path(__file__).resolve().parent
ASSETS_DIR = BACKEND_DIR / "assets"
OUTPUT_PATH = BACKEND_DIR / "exec_summary_template.docx"


def _patch_settings_update_fields(docx_path: Path) -> None:
    """Ensure word/settings.xml requests field refresh on open (TOC, PAGE). Matches exec_summary post-save patch."""
    try:
        buf = io.BytesIO()
        with zipfile.ZipFile(docx_path, "r") as zin:
            st = zin.read("word/settings.xml").decode("utf-8")
            if re.search(r"<w:updateFields\b", st):
                st = re.sub(
                    r"<w:updateFields[^>]*/>",
                    '<w:updateFields w:val="true"/>',
                    st,
                    count=1,
                )
            else:
                st = re.sub(
                    r"(<w:settings[^>]*>)",
                    r'\1<w:updateFields w:val="true"/>',
                    st,
                    count=1,
                )
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    if item.filename == "word/settings.xml":
                        data = st.encode("utf-8")
                    zout.writestr(item, data)
        docx_path.write_bytes(buf.getvalue())
    except OSError:
        pass


# Brand colors (match CSS Executive Summary template)
CSS_PURPLE = RGBColor(0x70, 0x30, 0xA0)
CSS_GREEN = RGBColor(0x00, 0x83, 0x3B)
CSS_RED = RGBColor(0x60, 0x00, 0x00)
CSS_BLUE = RGBColor(0x00, 0x37, 0x60)
CSS_BODY_GRAY = RGBColor(0x4D, 0x4D, 0x4D)
CSS_BLACK = RGBColor(0x00, 0x00, 0x00)


def _add_page_field(paragraph) -> None:
    """PAGE field: one OOXML element per run (same rule as TOC) so Word evaluates the field reliably."""
    r1 = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    r1._r.append(fld_begin)

    r2 = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    r2._r.append(instr)

    r3 = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    r3._r.append(fld_sep)

    r4 = paragraph.add_run()
    t = OxmlElement("w:t")
    t.text = "1"
    r4._r.append(t)

    r5 = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r5._r.append(fld_end)


def _define_styles(doc: Document) -> None:
    styles = doc.styles

    def add_para_style(name: str, font_name: str, size_pt: float, bold: bool, color: RGBColor | None, justify: bool = False):
        try:
            st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        except ValueError:
            st = styles[name]
        st.base_style = styles["Normal"]
        f = st.font
        f.name = font_name
        f.size = Pt(size_pt)
        f.bold = bold
        if color is not None:
            f.color.rgb = color
        if justify:
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Calibri fallback (Altone Trial / Raleway may not be installed)
    add_para_style("CSSHeading1", "Calibri", 24, True, CSS_PURPLE, False)
    add_para_style("CSSHeading2", "Calibri", 18, True, CSS_PURPLE, False)
    add_para_style("CSSHeading2Green", "Calibri", 18, True, CSS_GREEN, False)
    add_para_style("CSSHeading2Red", "Calibri", 18, True, CSS_RED, False)
    add_para_style("CSSHeading2Blue", "Calibri", 18, True, CSS_BLUE, False)
    add_para_style("CSSBody", "Calibri", 11, False, CSS_BODY_GRAY, True)
    add_para_style("CSSBodyBold", "Calibri", 11, True, CSS_BLACK, True)
    add_para_style("CSSConclusionTitle", "Calibri", 22, True, CSS_PURPLE, False)
    add_para_style("CSSTOCTitle", "Calibri", 20, True, CSS_PURPLE, False)
    add_para_style("CSSCoverTitle", "Calibri", 18, True, CSS_PURPLE, False)
    add_para_style("CSSCoverSub", "Calibri", 14, False, CSS_PURPLE, False)


def _override_builtin_heading_styles(doc: Document) -> None:
    """Match former CSSHeading1/2 visuals on Word built-in Heading 1 / Heading 2 (TOC-friendly)."""
    try:
        h1 = doc.styles["Heading 1"]
        h1.font.name = "Calibri"
        h1.font.size = Pt(24)
        h1.font.bold = True
        h1.font.color.rgb = CSS_PURPLE
    except KeyError:
        pass
    try:
        h2 = doc.styles["Heading 2"]
        h2.font.name = "Calibri"
        h2.font.size = Pt(18)
        h2.font.bold = True
        h2.font.color.rgb = CSS_PURPLE
    except KeyError:
        pass


def _clear_header_footer(section) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for hdr in (section.header, section.footer):
        while hdr.paragraphs:
            p0 = hdr.paragraphs[0]
            p0._element.getparent().remove(p0._element)
        hdr.add_paragraph()


def _build() -> None:
    cover = ASSETS_DIR / "cover_bg.png"
    hlogo = ASSETS_DIR / "header_logo.png"
    fbanner = ASSETS_DIR / "footer_banner.png"
    for p in (cover, hlogo, fbanner):
        if not p.is_file():
            raise FileNotFoundError(f"Missing asset: {p}")

    doc = Document()
    _define_styles(doc)
    _override_builtin_heading_styles(doc)

    sec0 = doc.sections[0]
    # Full-bleed cover (US Letter 8.5 × 11 in)
    sec0.left_margin = Inches(0)
    sec0.right_margin = Inches(0)
    sec0.top_margin = Inches(0)
    sec0.bottom_margin = Inches(0)

    _clear_header_footer(sec0)

    # Cover: background image fills page
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_img = p_img.add_run()
    run_img.add_picture(str(cover), width=Inches(8.5), height=Inches(11))

    doc.add_paragraph()
    p_org = doc.add_paragraph()
    p_org.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_org.style = doc.styles["CSSCoverTitle"]
    p_org.add_run("{{ORG_NAME}}")

    p_asmt = doc.add_paragraph()
    p_asmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_asmt.style = doc.styles["CSSCoverTitle"]
    p_asmt.add_run("{{ASSESSMENT_TYPE}}")

    p_client = doc.add_paragraph()
    p_client.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_client.style = doc.styles["CSSCoverSub"]
    p_client.add_run("For {{CLIENT_NAME}} Infrastructure")

    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date.style = doc.styles["CSSCoverSub"]
    p_date.add_run("{{REPORT_DATE}}")

    doc.add_paragraph()

    # Content section: header logo + footer banner + PAGE
    doc.add_section(WD_SECTION.NEW_PAGE)
    sec1 = doc.sections[1]
    sec1.left_margin = Inches(1)
    sec1.right_margin = Inches(1)
    sec1.top_margin = Inches(1)
    sec1.bottom_margin = Inches(1)
    sec1.header.is_linked_to_previous = False
    sec1.footer.is_linked_to_previous = False

    # Header: logo + optional doc id line (filled at generation)
    hp = sec1.header.paragraphs[0] if sec1.header.paragraphs else sec1.header.add_paragraph()
    hp.clear()
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run()
    hr.add_picture(str(hlogo), width=Inches(6.2))

    hp2 = sec1.header.add_paragraph()
    hp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp2.add_run("{{DOCUMENT_ID}}")

    # Footer: banner + page number
    fp = sec1.footer.paragraphs[0] if sec1.footer.paragraphs else sec1.footer.add_paragraph()
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run()
    fr.add_picture(str(fbanner), width=Inches(6.5))

    fp2 = sec1.footer.add_paragraph()
    fp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_page_field(fp2)

    # Placeholder paragraph — removed when generating full report (content appended after)
    p_start = doc.add_paragraph()
    p_start.style = doc.styles["CSSBody"]
    p_start.add_run("{{CONTENT_START}}")

    doc.save(str(OUTPUT_PATH))
    _patch_settings_update_fields(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH} (with w:updateFields in settings.xml)")


if __name__ == "__main__":
    _build()
