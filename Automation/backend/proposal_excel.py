"""
Reader for the CyberSmithSECURE "Pre-Engagement Scoping v2" workbook.

Turns an uploaded .xlsx into a ScopingResult: client details, reference number and the list of
in-scope annexures, each with a Count and a Quantity (docs/ADVANCED_FEATURE_PLAN.md, D30-D33):

    Count    = number of rows in the annexure's block holding a number above 0
    Quantity = sum of those numbers

All mapping (which cover cell, which matrix rows, labels) lives in annexure_map.json.
The reader never raises on a bad workbook: problems come back in ScopingResult.errors
(file rejected) or ScopingResult.warnings (processing continued).
"""
from __future__ import annotations

import datetime as dt
import json
import zipfile
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any

import openpyxl

HERE = Path(__file__).resolve().parent
MAP_FILE = HERE / "annexure_map.json"
ALLOWED_EXTENSIONS = {".xlsx"}


@lru_cache(maxsize=1)
def load_map() -> dict[str, Any]:
    with open(MAP_FILE, encoding="utf-8-sig") as fh:  # tolerates an editor-added BOM
        return json.load(fh)


@dataclass
class RowDetail:
    """One matrix row inside an annexure's block (shown on the review screen)."""
    row: int
    category: str
    value: Any
    status: str  # counted | blank | zero | text | negative | uncached_formula

    def to_dict(self) -> dict[str, Any]:
        return {"row": self.row, "category": self.category, "value": self.value, "status": self.status}


@dataclass
class AnnexureEntry:
    type_key: str
    type_label: str
    annexure: int | str            # template annexure number or "custom"
    catalog_key: str | None        # key in css_proposal_generator.ASSESSMENT_CATALOG
    count: int
    quantity: int | float
    count_label: str
    frequency: str
    rows: list[RowDetail] = field(default_factory=list)

    @property
    def needs_confirmation(self) -> bool:
        """D16/D33: the block holds nothing above 0, so the user must confirm."""
        return self.count == 0

    def to_dict(self) -> dict[str, Any]:
        return {
            "type": self.type_label,
            "annexure_no_or_custom": self.annexure,
            "count": self.count,
            "quantity": self.quantity,
            "count_label": self.count_label,
            "frequency": self.frequency,
            "type_key": self.type_key,
            "catalog_key": self.catalog_key,
            "needs_confirmation": self.needs_confirmation,
            "rows": [r.to_dict() for r in self.rows],
        }


@dataclass
class ScopingResult:
    client: dict[str, Any] = field(default_factory=dict)
    reference_no: str | None = None
    included: list[AnnexureEntry] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    # Answers from the other sheets, kept for the AI prompt only (plan D13); see extract_context().
    context: dict[str, Any] = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return not self.errors

    def distinct_annexures(self) -> list[int | str]:
        seen: list[int | str] = []
        for entry in self.included:
            if entry.annexure not in seen:
                seen.append(entry.annexure)
        return seen

    def pricing_lines(self) -> list[dict[str, Any]]:
        """One line per in-scope type (External and Internal Network PT stay separate, D11)."""
        return [
            {
                "label": e.type_label,
                "annexure_no_or_custom": e.annexure,
                "quantity": e.quantity,
                "frequency": e.frequency,
            }
            for e in self.included
        ]

    def to_dict(self) -> dict[str, Any]:
        return {
            "included_annexures": [e.to_dict() for e in self.included],
            "pricing_lines": self.pricing_lines(),
            "distinct_annexures": self.distinct_annexures(),
            "warnings": list(self.warnings),
            "client": dict(self.client),
            "reference_no": self.reference_no,
            "errors": list(self.errors),
        }


# --------------------------------------------------------------------------- helpers
def _is_yes(value: Any, yes_tokens: set[str]) -> bool:
    if value is None or isinstance(value, bool):  # a boolean TRUE is not an accepted token
        return False
    return str(value).strip().lower() in yes_tokens


def _as_number(value: Any) -> int | float | None:
    """Real Excel numbers only. Booleans and numeric-looking text are not numbers."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    return None


def _clean_number(n: int | float) -> int | float:
    if isinstance(n, float) and n.is_integer():
        return int(n)
    return n


def _read_client_field(ws, coord: str, field_name: str, cfg: dict[str, Any], warnings: list[str]):
    value = ws[coord].value
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    text = str(value).strip()
    if text == "" or text in cfg["placeholders"]:
        return None
    if field_name in cfg["date_fields"]:
        warnings.append(f"date_not_real_date: {ws.title}!{coord}={text!r}")
    return text


def _block_stats(ws_val, ws_formula, rows: list[int], cfg: dict[str, Any], warnings: list[str]):
    """Count and sum one block of matrix rows. Returns (count, quantity, row details)."""
    qty_col = cfg["matrix_quantity_column"]
    cat_col = cfg["matrix_category_column"]
    count = 0
    total: int | float = 0
    details: list[RowDetail] = []
    for r in rows:
        coord = f"{qty_col}{r}"
        value = ws_val[coord].value
        raw = ws_formula[coord].value
        category = str(ws_val[f"{cat_col}{r}"].value or "").strip()
        if value is None:
            if isinstance(raw, str) and raw.startswith("="):
                warnings.append(f"formula_no_cached_value: {ws_val.title}!{coord}")
                details.append(RowDetail(r, category, raw, "uncached_formula"))
            else:
                details.append(RowDetail(r, category, None, "blank"))
            continue
        number = _as_number(value)
        if number is None:
            warnings.append(f"non_numeric_quantity_skipped: {category} ({ws_val.title}!{coord})={value!r}")
            details.append(RowDetail(r, category, value, "text"))
            continue
        if number < 0:
            warnings.append(f"negative_quantity_skipped: {category} ({ws_val.title}!{coord})={value!r}")
            details.append(RowDetail(r, category, _clean_number(number), "negative"))
            continue
        if number == 0:
            details.append(RowDetail(r, category, 0, "zero"))
            continue
        count += 1
        total += number
        details.append(RowDetail(r, category, _clean_number(number), "counted"))
    return count, _clean_number(total), details


# --------------------------------------------------------------------------- main entry
def parse_scoping_workbook(path: str | Path) -> ScopingResult:
    """Read one scoping workbook. Never raises; see ScopingResult.errors / .warnings."""
    cfg = load_map()
    result = ScopingResult()
    path = Path(path)

    # -- file-level checks (before anything is opened) ----------------------------------
    if not path.is_file():
        result.errors.append("file_not_found")
        return result
    if path.suffix.lower() not in ALLOWED_EXTENSIONS:
        result.errors.append(f"invalid_extension: {path.suffix or '(none)'} (only .xlsx is accepted)")
        return result
    size = path.stat().st_size
    if size > cfg["max_file_bytes"]:
        result.errors.append(f"file_too_large: {size} bytes > {cfg['max_file_bytes']}")
        return result
    if not zipfile.is_zipfile(path):
        result.errors.append("invalid_xlsx: not a valid zip/OOXML container")
        return result
    try:
        with zipfile.ZipFile(path) as zf:
            if sum(info.file_size for info in zf.infolist()) > cfg["max_uncompressed_bytes"]:
                result.errors.append("archive_too_large: uncompressed content exceeds the limit")
                return result
    except zipfile.BadZipFile:
        result.errors.append("invalid_xlsx: not a valid zip/OOXML container")
        return result

    try:
        wb_val = openpyxl.load_workbook(path, data_only=True, keep_links=False)
        wb_formula = openpyxl.load_workbook(path, data_only=False, keep_links=False)
    except Exception as exc:  # truncated or malformed OOXML
        result.errors.append(f"corrupt_file: {type(exc).__name__}")
        return result

    # -- structure checks ---------------------------------------------------------------
    cover_name, matrix_name = cfg["sheets"]["cover"], cfg["sheets"]["matrix"]
    for name in (cover_name, matrix_name):
        if name not in wb_val.sheetnames:
            result.errors.append(f"missing_sheet: {name}")
    if result.errors:
        return result

    cover = wb_val[cover_name]
    matrix_val = wb_val[matrix_name]
    matrix_formula = wb_formula[matrix_name]

    header = cfg["matrix_quantity_header"]
    found = matrix_val[header["cell"]].value
    if str(found).strip() != header["text"]:
        result.errors.append(
            f"quantity_header_changed: expected {header['text']!r} at {matrix_name}!{header['cell']}, "
            f"found {found!r}"
        )
        return result

    # -- client block and reference number ----------------------------------------------
    for field_name, coord in cfg["client_cells"].items():
        result.client[field_name] = _read_client_field(cover, coord, field_name, cfg, result.warnings)

    ref_raw = cover[cfg["reference_cell"]].value
    if ref_raw is not None:
        ref = str(ref_raw).strip()
        if ref and ref not in cfg["placeholders"] and ref != cfg["reference_placeholder"]:
            result.reference_no = ref
    if result.reference_no is None:
        result.warnings.append("reference_required")

    # -- which types are in scope -------------------------------------------------------
    yes = set(cfg["yes_tokens"])
    all_types = cfg["types"] + cfg["types_without_matrix_rows"]
    in_scope = {t["key"]: _is_yes(cover[t["cover_cell"]].value, yes) for t in all_types}
    if not any(in_scope.values()):
        result.errors.append("no_assessment_types_in_scope")

    frequency = cfg["default_frequency"]
    for t in cfg["types"]:
        if not in_scope[t["key"]]:
            continue
        count, quantity, details = _block_stats(matrix_val, matrix_formula, t["rows"], cfg, result.warnings)
        entry = AnnexureEntry(t["key"], t["label"], t["annexure"], t["catalog_key"],
                              count, quantity, t["count_label"], frequency, details)
        result.included.append(entry)
        if entry.needs_confirmation:
            result.warnings.append(f"quantity_zero_confirm: {t['label']}")

    # Types with no matrix rows: Count 1 / Quantity 1, edited by the user in the preview.
    for t in cfg["types_without_matrix_rows"]:
        if in_scope[t["key"]]:
            result.included.append(AnnexureEntry(
                t["key"], t["label"], t["annexure"], t["catalog_key"],
                t["default_count"], t["default_quantity"], t["count_label"], frequency))

    try:
        result.context = extract_context(wb_val, cfg)
    except Exception as exc:  # context is a bonus for the AI; it must never fail the upload
        result.warnings.append(f"context_unreadable: {type(exc).__name__}")
    return result


# --------------------------------------------------------------------------- AI context (D13)
def _clip(value: Any, limit: int) -> str:
    text = " ".join(str(value).split()) if value is not None else ""
    return text if len(text) <= limit else text[: limit - 1] + "…"


def _qa_rows(ws, rows: tuple[int, int], q_col: str, a_cols: list[str], limit: int) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for r in range(rows[0], rows[1] + 1):
        question = ws[f"{q_col}{r}"].value
        if question is None or not str(question).strip():
            continue
        answers = [_clip(ws[f"{c}{r}"].value, limit) for c in a_cols]
        answers = [a for a in answers if a]
        if answers:                       # unanswered questions carry no information
            items.append({"question": _clip(question, 200), "answers": answers})
    return items


def extract_context(wb, cfg: dict[str, Any] | None = None) -> dict[str, Any]:
    """Read the answered questions of the non-driving sheets.

    Returns {"general": {section: [{question, answers}]}, "detail": {key: [...]},
             "roe": {"parameters": [...], "compliance": [...]}}.
    Rules of Engagement contacts are personal data and are not read at all.
    """
    cfg = cfg or load_map()
    ctx = cfg["context"]
    limit = ctx["max_answer_chars"]
    out: dict[str, Any] = {"general": {}, "detail": {}, "roe": {"parameters": [], "compliance": []}}

    if ctx["general_sheet"] in wb.sheetnames:
        ws = wb[ctx["general_sheet"]]
        for section, rows in ctx["general_sections"].items():
            out["general"][section] = _qa_rows(ws, tuple(rows), ctx["general_question_column"],
                                               ctx["general_answer_columns"], limit)
    for key, spec in ctx["detail_sheets"].items():
        if spec["sheet"] in wb.sheetnames:
            out["detail"][key] = _qa_rows(wb[spec["sheet"]], tuple(spec["rows"]), spec["question_column"],
                                          spec["answer_columns"], limit)
    roe = ctx["roe"]
    if roe["sheet"] in wb.sheetnames:
        ws = wb[roe["sheet"]]
        out["roe"]["parameters"] = _qa_rows(ws, tuple(roe["parameter_rows"]), roe["label_column"],
                                            [roe["answer_column"]], limit)
        out["roe"]["compliance"] = _qa_rows(ws, tuple(roe["compliance_rows"]), roe["label_column"],
                                            [roe["answer_column"]], limit)
    return out


def context_block_for_type(context: dict[str, Any], type_key: str, cfg: dict[str, Any] | None = None) -> str:
    """Plain-text block of the answers relevant to one assessment type (for the AI prompt)."""
    cfg = cfg or load_map()
    ctx_cfg = cfg["context"]
    sources = ctx_cfg["type_sources"].get(type_key, {})
    lines: list[str] = []

    def add(title: str, items: list[dict[str, Any]]) -> None:
        if not items:
            return
        lines.append(f"{title}:")
        for it in items:
            lines.append(f"- {it['question']} => {' | '.join(it['answers'])}")

    if sources.get("general"):
        add("Scoping answers", context.get("general", {}).get(sources["general"], []))
    if sources.get("detail"):
        add("Detailed information gathering", context.get("detail", {}).get(sources["detail"], []))
    add("Rules of engagement", context.get("roe", {}).get("parameters", []))
    add("Compliance requirements", context.get("roe", {}).get("compliance", []))
    block = "\n".join(lines)
    limit = ctx_cfg["max_block_chars"]
    return block if len(block) <= limit else block[: limit - 1] + "…"


def suggest_sector(regulatory_driver: str | None, cfg: dict[str, Any] | None = None) -> str:
    """D20: a suggestion only (e.g. RBI/SEBI -> BFSI); the user always confirms."""
    cfg = cfg or load_map()
    text = (regulatory_driver or "").lower()
    for token, sector in cfg["sector_suggestions"].items():
        if token in text:
            return sector
    return ""
