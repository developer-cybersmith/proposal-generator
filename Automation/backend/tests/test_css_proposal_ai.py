"""
Unit tests for the AI-powered custom annexure generation feature.

Run from Automation/backend/:
    python -m pytest tests/test_css_proposal_ai.py -v

Or directly:
    python tests/test_css_proposal_ai.py
"""
from __future__ import annotations

import sys
import os
import json
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

# Ensure the backend package is importable
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from css_proposal_generator import (
    _parse_annexure_line,
    validate_ai_annexure_output,
    _parse_ai_context,
    _validate_ai_context,
    _build_ai_user_message,
    _make_request_catalog,
    _register_custom_annexure,
    _toc_title_for_key,
    _refresh_all_toc_pages,
    _update_contents_toc_structure,
    _compute_layout_pages,
    _find_contents_region,
    _insert_logo_on_cover,
    _find_cover_logo_paragraph,
    build_custom_annexure_block,
    BASE_TEMPLATE,
    CSS_HEADING_COLOR,
    CSS_BODY_COLOR,
    OTHER_KEY_RE,
    AI_REQUIRED_SECTIONS,
    CSS_FONT_NAME,
    CSS_ANNEXURE_HEADING_SIZE,
    CSS_SECTION_LABEL_SIZE,
    CSS_BODY_SIZE,
    MASTER_BD_PROMPT_FILE,
)


# ---------------------------------------------------------------------------
# 1. Line parser tests
# ---------------------------------------------------------------------------

class TestParseAnnexureLine(unittest.TestCase):

    def test_blank_line(self):
        result = _parse_annexure_line("")
        self.assertEqual(result["style"], "blank")
        self.assertEqual(result["text"], "")

    def test_blank_whitespace_line(self):
        result = _parse_annexure_line("   ")
        self.assertEqual(result["style"], "blank")

    def test_annexure_heading(self):
        result = _parse_annexure_line("Annexure 1: Network Penetration Testing Assessment")
        self.assertEqual(result["style"], "annexure_heading")
        self.assertIn("Network", result["text"])

    def test_annexure_heading_case_insensitive(self):
        result = _parse_annexure_line("annexure 3: GCP Cloud Security Review")
        self.assertEqual(result["style"], "annexure_heading")

    def test_section_label_standard(self):
        result = _parse_annexure_line("Sub-Activity:")
        self.assertEqual(result["style"], "section_label")

    def test_section_label_methodologies(self):
        result = _parse_annexure_line("Methodologies:")
        self.assertEqual(result["style"], "section_label")

    def test_section_label_deliverables(self):
        result = _parse_annexure_line("Deliverables:")
        self.assertEqual(result["style"], "section_label")

    def test_section_label_objectives(self):
        result = _parse_annexure_line("Objectives:")
        self.assertEqual(result["style"], "section_label")

    def test_main_bullet(self):
        result = _parse_annexure_line("\u2022 Review IAM policies and role assignments")
        self.assertEqual(result["style"], "main_bullet")
        self.assertIn("Review IAM", result["text"])

    def test_sub_bullet(self):
        result = _parse_annexure_line("o Check for privilege escalation paths")
        self.assertEqual(result["style"], "sub_bullet")
        self.assertIn("Check for", result["text"])

    def test_sub_bullet_not_of_on(self):
        # Lines starting with "of " or "on " should be body, not sub_bullet
        result_of = _parse_annexure_line("of the assessment")
        result_on = _parse_annexure_line("on the server")
        self.assertEqual(result_of["style"], "body")
        self.assertEqual(result_on["style"], "body")

    def test_body_text(self):
        result = _parse_annexure_line("This assessment covers all major attack surfaces.")
        self.assertEqual(result["style"], "body")
        self.assertIn("attack surfaces", result["text"])


# ---------------------------------------------------------------------------
# 2. AI output validation tests
# ---------------------------------------------------------------------------

VALID_AI_OUTPUT = """
Annexure 5: Cloud Red Teaming Assessment

Sub-Activity:
\u2022 Simulate adversary attack paths on cloud infrastructure.
o Identify privilege escalation paths.
o Test lateral movement between cloud services.

Methodologies:
\u2022 MITRE ATT&CK Cloud Matrix
\u2022 Manual exploitation of misconfigured IAM roles

Cloud Red Teaming Assessment:
\u2022 Attack surface enumeration
\u2022 Identity and Access Management review

Key Phases:
\u2022 Reconnaissance
\u2022 Initial Access
\u2022 Lateral Movement

Frequency:
\u2022 One-time engagement

Count:
\u2022 2 cloud accounts (AWS + Azure)

Deliverables:
\u2022 Executive Summary Report
\u2022 Technical Findings Report

Expectation:
\u2022 Detailed remediation guidance with risk ratings.

Objectives:
\u2022 Identify and validate all critical cloud attack paths.

Out of Scope/Exclusions:
\u2022 Physical security testing
\u2022 On-premises infrastructure

Requirements for Starting:
\u2022 Cloud read-only access credentials
\u2022 Architecture diagrams

Comments:
\u2022 Report will be delivered within 5 business days.
""".strip()


class TestValidateAIAnnexureOutput(unittest.TestCase):

    def test_valid_output_passes(self):
        validate_ai_annexure_output(VALID_AI_OUTPUT, "Cloud Red Teaming")  # no raise

    def test_too_short_raises(self):
        with self.assertRaises(ValueError) as ctx:
            validate_ai_annexure_output("Short text here.", "Test")
        # Validator raises for missing sections first, or too-short — either is acceptable
        msg = str(ctx.exception).lower()
        self.assertTrue(
            "too short" in msg or "missing" in msg,
            f"Expected 'too short' or 'missing' in: {msg}"
        )

    def test_missing_section_raises(self):
        truncated = VALID_AI_OUTPUT.replace("Sub-Activity:", "")
        with self.assertRaises(ValueError) as ctx:
            validate_ai_annexure_output(truncated, "Cloud Red Teaming")
        self.assertIn("missing", str(ctx.exception).lower())

    def test_all_required_sections_present(self):
        for section in AI_REQUIRED_SECTIONS:
            # Ensure VALID_AI_OUTPUT actually contains each required section
            self.assertIn(
                section.lower(), VALID_AI_OUTPUT.lower(),
                msg=f"Test fixture missing: {section}"
            )


# ---------------------------------------------------------------------------
# 3. AIContextSchema parsing & validation tests
# ---------------------------------------------------------------------------

VALID_AI_CTX_DICT = {
    "sector": "Insurance",
    "engagement_summary": "Comprehensive cloud and endpoint security assessment.",
    "primary_tier": "Tier 3 \u2013 Advanced",
    "tier_notes": "",
    "counts": {
        "ad_identity": "2 AD domains",
        "cloud": "AWS",
        "apps_apis": "",
        "infra": "50 IPs",
        "users": "200 employees",
        "duration": "4 weeks",
    },
    "commercials": {
        "total_cost": "4,33,000 INR",
        "taxes": "18% GST",
    },
    "payment_model": "hybrid_v2",
    "special_constraints": "",
    "custom_annexures": [
        {
            "slot_index": 0,
            "key": "__other__1234567890_abc12",
            "title": "Cloud Red Teaming Assessment",
            "scope": "Simulate adversary attack on AWS + Azure. Identify lateral movement paths.",
        }
    ],
}


class TestAIContextSchema(unittest.TestCase):

    def test_parse_valid_context(self):
        ctx = _parse_ai_context(json.dumps(VALID_AI_CTX_DICT))
        self.assertEqual(ctx.sector, "Insurance")
        self.assertEqual(ctx.counts.duration, "4 weeks")
        self.assertEqual(ctx.commercials.total_cost, "4,33,000 INR")
        self.assertEqual(ctx.payment_model, "hybrid_v2")
        self.assertEqual(len(ctx.custom_annexures), 1)
        self.assertEqual(ctx.custom_annexures[0].title, "Cloud Red Teaming Assessment")

    def test_parse_empty_string_returns_default(self):
        ctx = _parse_ai_context("")
        self.assertIsNotNone(ctx)

    def test_validate_passes_for_valid_context(self):
        ctx = _parse_ai_context(json.dumps(VALID_AI_CTX_DICT))
        _validate_ai_context(ctx)  # no raise

    def test_validate_fails_missing_sector(self):
        data = dict(VALID_AI_CTX_DICT)
        data["sector"] = ""
        ctx = _parse_ai_context(json.dumps(data))
        with self.assertRaises(ValueError) as cm:
            _validate_ai_context(ctx)
        self.assertIn("sector", str(cm.exception).lower())

    def test_validate_fails_short_engagement_summary(self):
        data = {**VALID_AI_CTX_DICT, "engagement_summary": "Too short"}
        ctx = _parse_ai_context(json.dumps(data))
        with self.assertRaises(ValueError) as cm:
            _validate_ai_context(ctx)
        self.assertIn("engagement_summary", str(cm.exception).lower())

    def test_validate_fails_missing_duration(self):
        data = json.loads(json.dumps(VALID_AI_CTX_DICT))
        data["counts"]["duration"] = ""
        ctx = _parse_ai_context(json.dumps(data))
        with self.assertRaises(ValueError) as cm:
            _validate_ai_context(ctx)
        self.assertIn("duration", str(cm.exception).lower())

    def test_validate_fails_missing_total_cost(self):
        data = json.loads(json.dumps(VALID_AI_CTX_DICT))
        data["commercials"]["total_cost"] = ""
        ctx = _parse_ai_context(json.dumps(data))
        with self.assertRaises(ValueError) as cm:
            _validate_ai_context(ctx)
        self.assertIn("total_cost", str(cm.exception).lower())

    def test_validate_fails_invalid_payment_model(self):
        data = {**VALID_AI_CTX_DICT, "payment_model": "invalid_model"}
        ctx = _parse_ai_context(json.dumps(data))
        with self.assertRaises(ValueError) as cm:
            _validate_ai_context(ctx)
        self.assertIn("payment_model", str(cm.exception).lower())

    def test_validate_fails_short_scope(self):
        data = json.loads(json.dumps(VALID_AI_CTX_DICT))
        data["custom_annexures"][0]["scope"] = "Too short"
        ctx = _parse_ai_context(json.dumps(data))
        with self.assertRaises(ValueError) as cm:
            _validate_ai_context(ctx)
        self.assertIn("scope", str(cm.exception).lower())

    def test_validate_fails_empty_title(self):
        data = json.loads(json.dumps(VALID_AI_CTX_DICT))
        data["custom_annexures"][0]["title"] = ""
        ctx = _parse_ai_context(json.dumps(data))
        with self.assertRaises(ValueError) as cm:
            _validate_ai_context(ctx)
        self.assertIn("title", str(cm.exception).lower())


# ---------------------------------------------------------------------------
# 4. Catalog mutation tests
# ---------------------------------------------------------------------------

class TestRequestCatalog(unittest.TestCase):

    def test_make_request_catalog_returns_copy(self):
        catalog = _make_request_catalog()
        self.assertIsInstance(catalog, dict)
        self.assertIn("post_breach", catalog)  # known static key

    def test_register_custom_annexure(self):
        catalog = _make_request_catalog()
        key = "__other__9999999999_test1"
        _register_custom_annexure(
            catalog=catalog,
            key=key,
            title="Test Custom Service",
            toc_title="Test Custom Service",
            slot=5,
            ai_elements=[],
        )
        self.assertIn(key, catalog)
        self.assertTrue(catalog[key]["is_custom"])
        self.assertEqual(catalog[key]["title"], "Test Custom Service")
        self.assertEqual(catalog[key]["custom_slot"], 5)

    def test_register_does_not_mutate_base(self):
        from css_proposal_generator import _BASE_CATALOG_BY_KEY
        catalog = _make_request_catalog()
        _register_custom_annexure(
            catalog=catalog,
            key="__other__9999_qqqqq",
            title="Mutate Test",
            toc_title="Mutate Test",
            slot=1,
        )
        self.assertNotIn("__other__9999_qqqqq", _BASE_CATALOG_BY_KEY)


# ---------------------------------------------------------------------------
# 5. Other-key regex tests
# ---------------------------------------------------------------------------

class TestOtherKeyRe(unittest.TestCase):

    def test_valid_key(self):
        self.assertTrue(OTHER_KEY_RE.match("__other__1716800000000_ab3de"))

    def test_invalid_key_no_ts(self):
        self.assertIsNone(OTHER_KEY_RE.match("__other__"))

    def test_invalid_key_normal_catalog_key(self):
        self.assertIsNone(OTHER_KEY_RE.match("post_breach"))


# ---------------------------------------------------------------------------
# 6. Brand constants smoke test
# ---------------------------------------------------------------------------

class TestBrandConstants(unittest.TestCase):

    def test_font_name(self):
        self.assertEqual(CSS_FONT_NAME, "Corbel")

    def test_heading_size(self):
        from docx.shared import Pt
        self.assertEqual(CSS_ANNEXURE_HEADING_SIZE, Pt(16))

    def test_section_label_size(self):
        from docx.shared import Pt
        self.assertEqual(CSS_SECTION_LABEL_SIZE, Pt(14))

    def test_body_size(self):
        from docx.shared import Pt
        self.assertEqual(CSS_BODY_SIZE, Pt(11))

    def test_heading_color_matches_static_template(self):
        self.assertEqual(str(CSS_HEADING_COLOR), "7030A0")

    def test_body_color_is_black(self):
        self.assertEqual(str(CSS_BODY_COLOR), "000000")


# ---------------------------------------------------------------------------
# 6b. Regression: AI annexure heading/label colors match static template
# ---------------------------------------------------------------------------

class TestAnnexureColorRegression(unittest.TestCase):

    SAMPLE_AI_TEXT = """Annexure 1: Cloud Red Teaming Assessment

Sub-Activity:
\u2022 Simulate adversary attack paths on cloud infrastructure.

Methodologies:
\u2022 MITRE ATT&CK for Cloud
"""

    def _run_rgb(self, run) -> str | None:
        rgb = run.font.color.rgb
        return str(rgb) if rgb is not None else None

    def test_build_custom_annexure_heading_is_purple(self):
        from docx import Document
        from docx.oxml.ns import qn
        import copy

        elements = build_custom_annexure_block(
            raw_ai_text=self.SAMPLE_AI_TEXT,
            slot_number=1,
            title="Cloud Red Teaming Assessment",
            client_name="Test Corp",
        )
        doc = Document()
        for p in list(doc.paragraphs):
            p._element.getparent().remove(p._element)
        body = doc.element.body
        for el in elements:
            if el.tag != qn("w:sectPr"):
                body.append(copy.deepcopy(el))

        heading_para = next(
            (p for p in doc.paragraphs if (p.text or "").startswith("Annexure 1:")),
            None,
        )
        self.assertIsNotNone(heading_para)
        self.assertEqual(self._run_rgb(heading_para.runs[0]), "7030A0")
        self.assertTrue(heading_para.runs[0].font.bold)
        self.assertTrue(heading_para.runs[0].font.italic)

    def test_build_custom_annexure_section_label_is_purple(self):
        from docx import Document
        from docx.oxml.ns import qn
        import copy

        elements = build_custom_annexure_block(
            raw_ai_text=self.SAMPLE_AI_TEXT,
            slot_number=1,
            title="Cloud Red Teaming Assessment",
            client_name="Test Corp",
        )
        doc = Document()
        for p in list(doc.paragraphs):
            p._element.getparent().remove(p._element)
        body = doc.element.body
        for el in elements:
            if el.tag != qn("w:sectPr"):
                body.append(copy.deepcopy(el))

        label_para = next(
            (p for p in doc.paragraphs if (p.text or "").strip() == "Sub-Activity:"),
            None,
        )
        self.assertIsNotNone(label_para)
        self.assertEqual(self._run_rgb(label_para.runs[0]), "7030A0")
        self.assertTrue(label_para.runs[0].font.bold)
        self.assertNotEqual(label_para.runs[0].font.italic, True)

    def test_build_custom_annexure_body_bullet_is_black(self):
        from docx import Document
        from docx.oxml.ns import qn
        import copy

        elements = build_custom_annexure_block(
            raw_ai_text=self.SAMPLE_AI_TEXT,
            slot_number=1,
            title="Cloud Red Teaming Assessment",
            client_name="Test Corp",
        )
        doc = Document()
        for p in list(doc.paragraphs):
            p._element.getparent().remove(p._element)
        body = doc.element.body
        for el in elements:
            if el.tag != qn("w:sectPr"):
                body.append(copy.deepcopy(el))

        bullet_para = next(
            (p for p in doc.paragraphs if "Simulate adversary" in (p.text or "")),
            None,
        )
        self.assertIsNotNone(bullet_para)
        self.assertEqual(self._run_rgb(bullet_para.runs[0]), "000000")


# ---------------------------------------------------------------------------
# 7. User message builder smoke test
# ---------------------------------------------------------------------------

class TestBuildAIUserMessage(unittest.TestCase):

    def _make_context_and_ca(self):
        ctx = _parse_ai_context(json.dumps(VALID_AI_CTX_DICT))
        ca = ctx.custom_annexures[0]
        return ctx, ca

    def test_message_contains_company_name(self):
        ctx, ca = self._make_context_and_ca()
        msg = _build_ai_user_message("Acme Corp", ctx, ca, slot_number=5)
        self.assertIn("Acme Corp", msg)

    def test_message_contains_sector(self):
        ctx, ca = self._make_context_and_ca()
        msg = _build_ai_user_message("Acme Corp", ctx, ca, slot_number=5)
        self.assertIn("Insurance", msg)

    def test_message_contains_title(self):
        ctx, ca = self._make_context_and_ca()
        msg = _build_ai_user_message("Acme Corp", ctx, ca, slot_number=5)
        self.assertIn("Cloud Red Teaming", msg)

    def test_message_contains_do_not_generate_tc(self):
        ctx, ca = self._make_context_and_ca()
        msg = _build_ai_user_message("Acme Corp", ctx, ca, slot_number=5)
        self.assertIn("Do NOT generate the Terms", msg)

    def test_message_contains_output_instruction(self):
        ctx, ca = self._make_context_and_ca()
        msg = _build_ai_user_message("Acme Corp", ctx, ca, slot_number=5)
        self.assertIn("IMPORTANT OUTPUT INSTRUCTION", msg)


# ---------------------------------------------------------------------------
# 8. Master BD Prompt file test
# ---------------------------------------------------------------------------

class TestMasterBDPromptFile(unittest.TestCase):

    def test_prompt_file_exists(self):
        self.assertTrue(
            MASTER_BD_PROMPT_FILE.is_file(),
            f"master_bd_prompt.txt not found at {MASTER_BD_PROMPT_FILE}"
        )

    def test_prompt_file_not_empty(self):
        if MASTER_BD_PROMPT_FILE.is_file():
            content = MASTER_BD_PROMPT_FILE.read_text(encoding="utf-8")
            self.assertGreater(len(content), 500, "master_bd_prompt.txt seems too short")

    def test_prompt_file_has_annexure_layout(self):
        if MASTER_BD_PROMPT_FILE.is_file():
            content = MASTER_BD_PROMPT_FILE.read_text(encoding="utf-8")
            self.assertIn("Annexure", content)

    def test_prompt_file_no_my_input_section(self):
        """The example 'My Input' section must be stripped from the prompt file."""
        if MASTER_BD_PROMPT_FILE.is_file():
            content = MASTER_BD_PROMPT_FILE.read_text(encoding="utf-8")
            self.assertNotIn("My Input-", content, "Example 'My Input' section was not stripped")


# ---------------------------------------------------------------------------
# 9. Regression: TOC must show human title, not __other__ key
# ---------------------------------------------------------------------------

class TestTocCustomAnnexureRegression(unittest.TestCase):
    """Regression for finalize overwriting TOC with raw __other__ internal keys."""

    CUSTOM_KEY = "__other__9999999999_regtest"
    CUSTOM_TITLE = "Cloud Red Teaming Assessment"

    def _build_custom_catalog(self) -> dict:
        catalog = _make_request_catalog()
        _register_custom_annexure(
            catalog=catalog,
            key=self.CUSTOM_KEY,
            title=self.CUSTOM_TITLE,
            toc_title=self.CUSTOM_TITLE,
            slot=1,
            ai_elements=[],
        )
        return catalog

    def test_toc_title_for_key_uses_request_catalog(self):
        catalog = self._build_custom_catalog()
        title = _toc_title_for_key(self.CUSTOM_KEY, catalog)
        self.assertEqual(title, self.CUSTOM_TITLE)
        self.assertNotIn("__other__", title)

    def test_toc_title_for_key_without_catalog_falls_back_to_raw_key(self):
        """Without request catalog, unknown keys still fall back (documents the gap)."""
        title = _toc_title_for_key(self.CUSTOM_KEY, catalog=None)
        self.assertEqual(title, self.CUSTOM_KEY)

    def test_refresh_all_toc_pages_writes_human_title_for_custom_annexure(self):
        if not BASE_TEMPLATE.is_file():
            self.skipTest("CSS_Proposal_Base_Template.docx not available")

        from docx import Document

        catalog = self._build_custom_catalog()
        ordered_keys = [self.CUSTOM_KEY]

        doc = Document(str(BASE_TEMPLATE))
        _update_contents_toc_structure(doc, ordered_keys, catalog=catalog)
        pages = _compute_layout_pages(doc)
        _refresh_all_toc_pages(doc, ordered_keys, pages, catalog=catalog)

        _, first_assessment_idx, _, _ = _find_contents_region(doc)
        self.assertIsNotNone(first_assessment_idx)
        toc_line = (doc.paragraphs[first_assessment_idx].text or "").strip()
        self.assertIn(self.CUSTOM_TITLE, toc_line)
        self.assertNotIn("__other__", toc_line)


# ---------------------------------------------------------------------------
# 10. Regression: client logo left-aligned under "For"
# ---------------------------------------------------------------------------

class TestLogoAlignmentRegression(unittest.TestCase):

    def test_insert_logo_on_cover_uses_left_alignment(self):
        if not BASE_TEMPLATE.is_file():
            self.skipTest("CSS_Proposal_Base_Template.docx not available")

        try:
            from PIL import Image
        except ImportError:
            self.skipTest("Pillow not installed")

        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        import tempfile

        doc = Document(str(BASE_TEMPLATE))
        logo_para = _find_cover_logo_paragraph(doc)
        self.assertIsNotNone(logo_para, "Cover logo paragraph not found in template")

        with tempfile.TemporaryDirectory() as tmp:
            logo_path = Path(tmp) / "test_logo.png"
            Image.new("RGB", (120, 60), color=(30, 144, 255)).save(logo_path)

            self.assertTrue(_insert_logo_on_cover(doc, logo_path))
            self.assertEqual(logo_para.alignment, WD_ALIGN_PARAGRAPH.LEFT)

    def test_client_logo_placeholder_paragraph_is_left_in_template_after_build(self):
        """Template logo slot should be left-aligned (not center)."""
        if not BASE_TEMPLATE.is_file():
            self.skipTest("CSS_Proposal_Base_Template.docx not available")

        from css_proposal_generator import _ensure_client_logo_left_alignment
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        _ensure_client_logo_left_alignment(BASE_TEMPLATE)

        doc = Document(str(BASE_TEMPLATE))
        logo_para = _find_cover_logo_paragraph(doc)
        if logo_para is None:
            self.skipTest("No client logo paragraph in template")

        self.assertEqual(
            logo_para.alignment,
            WD_ALIGN_PARAGRAPH.LEFT,
            "Template client_logo paragraph should be left-aligned under 'For'",
        )


if __name__ == "__main__":
    unittest.main(verbosity=2)
