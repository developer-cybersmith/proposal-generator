"""
Tests for proposal_fields.py, proposal_pricing.py, proposal_lint.py, proposal_tokens.py and the
page-break fixes in css_proposal_generator.py.

Run from Automation/backend/:   python -m pytest tests/test_proposal_fields_pricing_lint.py -v
"""
from __future__ import annotations

import sys
from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

import proposal_fields as pf  # noqa: E402
import proposal_lint as lint  # noqa: E402
import proposal_pricing as pp  # noqa: E402
import proposal_tokens as tokens  # noqa: E402
import css_proposal_generator as gen  # noqa: E402


def _para(text: str):
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    p.append(r)
    return p


def _texts(elements):
    return [pf.paragraph_text(e) for e in elements if e.tag == qn("w:p")]


# ----------------------------------------------------------------------------- Count / Frequency
def test_count_line_text_joins_entries():
    assert pf.count_line_text([(6, "Web Application categories")]) == "6 Web Application categories"
    assert pf.count_line_text([(6, "External Infrastructure categories"), (12, "Internal Infrastructure categories")]) \
        == "6 External Infrastructure categories, 12 Internal Infrastructure categories"


def test_existing_lines_are_rewritten_and_formatted_from_tokens():
    block = [_para("Annexure 1: X"), _para("Frequency: One time with Revalidation"), _para("Count: API = 2350")]
    out = pf.apply_count_and_frequency(block, count_text="6 Web Application categories", frequency="One Time")
    assert out == {"count": "rewritten", "frequency": "kept"}
    texts = _texts(block)
    assert texts[1] == "Frequency: One time with Revalidation"       # template wording kept (default)
    assert texts[2] == "Count: 6 Web Application categories"
    runs = list(block[2].iter(qn("w:r")))
    assert len(runs) == 2                                              # bold label + regular value
    label_rpr, value_rpr = (r.find(qn("w:rPr")) for r in runs)
    assert label_rpr.find(qn("w:b")) is not None and value_rpr.find(qn("w:b")) is None
    assert label_rpr.find(qn("w:color")).get(qn("w:val")) == tokens.BRAND_PURPLE
    assert value_rpr.find(qn("w:color")).get(qn("w:val")) == tokens.BODY_TEXT
    assert label_rpr.find(qn("w:rFonts")).get(qn("w:ascii")) == tokens.FONT


def test_edited_frequency_replaces_the_template_wording():
    block = [_para("Frequency: One time with Revalidation"), _para("Count:")]
    pf.apply_count_and_frequency(block, count_text="3 SaaS Application categories", frequency="Quarterly")
    assert _texts(block) == ["Frequency: Quarterly", "Count: 3 SaaS Application categories"]


def test_missing_lines_are_inserted_before_deliverables():
    block = [_para("Annexure 7: Mobile"), _para("Methodologies:"), _para("Deliverables:"), _para("Report")]
    out = pf.apply_count_and_frequency(block, count_text="2 Mobile Application categories", frequency="One Time")
    assert out == {"count": "inserted", "frequency": "inserted"}
    texts = _texts(block)
    assert texts.index("Frequency: One Time") < texts.index("Count: 2 Mobile Application categories") < texts.index("Deliverables:")


def test_count_inserted_after_existing_frequency():
    block = [_para("Frequency: One Time with Revalidation"), _para("Deliverables:")]
    pf.apply_count_and_frequency(block, count_text="1 Engagement categories", frequency=None)
    assert _texts(block)[1] == "Count: 1 Engagement categories"


def test_manual_flow_leaves_lines_alone_when_nothing_is_requested():
    block = [_para("Frequency: One time"), _para("Count: 1")]
    out = pf.apply_count_and_frequency(block, count_text=None, frequency=None)
    # no count requested: frequency line gets spec formatting only, Count untouched
    assert out["count"] == "skipped"
    assert _texts(block)[1] == "Count: 1"


def test_cross_references_are_renumbered_but_headings_are_not():
    block = [_para("Annexure 4: Web"), _para("See Annexure 4 and Annexure 13 for details."), _para("Contents\tAnnexure 4")]
    changed = pf.renumber_cross_references(block, {4: 1, 13: 2})
    assert changed == 2
    assert _texts(block) == ["Annexure 4: Web", "See Annexure 1 and Annexure 2 for details.", "Contents\tAnnexure 4"]


# ----------------------------------------------------------------------------- pricing
@pytest.mark.parametrize("amount, expected", [
    (0, "₹ 0"), (999, "₹ 999"), (1000, "₹ 1,000"), (100000, "₹ 1,00,000"),
    (125000, "₹ 1,25,000"), (12345678, "₹ 1,23,45,678"), (Decimal("1234.5"), "₹ 1,234.50"),
])
def test_format_inr_uses_indian_grouping(amount, expected):
    assert pp.format_inr(amount) == expected


@pytest.mark.parametrize("text, expected", [
    ("50000", Decimal("50000")), ("₹ 1,25,000 /-", Decimal("125000")), ("Rs. 75,000", Decimal("75000")),
    ("", None), (None, None), ("TBD", None), ("-5", None), ("1e3x", None),
])
def test_parse_amount(text, expected):
    assert pp.parse_amount(text) == expected


def test_total_is_shown_only_when_every_row_is_priced():
    rows = [{"cost": "50000"}, {"cost": "1,25,000"}]
    assert pp.total_text(rows) == "₹ 1,75,000 /- * + 18% GST"
    assert pp.total_text(rows + [{"cost": ""}]) == tokens.PRICING_TOTAL_EMPTY_TEXT   # never a partial sum
    assert pp.total_text([{"cost": ""}]) == "₹ /- * + 18% GST"
    assert pp.total_text([]) == tokens.PRICING_TOTAL_EMPTY_TEXT


def test_pricing_table_structure_and_tokens():
    rows = [{"title": "Web Application Penetration Testing", "quantity": "23", "frequency": "One Time", "cost": ""},
            {"title": "Cloud Penetration Testing", "quantity": "5", "frequency": "One Time", "cost": ""}]
    width = 10914
    tbl = pp.build_pricing_table(rows, width)
    trs = tbl.findall(qn("w:tr"))
    assert len(trs) == 1 + len(rows) + 1                                   # header + rows + total
    grid = [int(g.get(qn("w:w"))) for g in tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol"))]
    assert sum(grid) == width and len(grid) == 5
    header_texts = ["".join(t.text for t in tc.iter(qn("w:t"))) for tc in trs[0].findall(qn("w:tc"))]
    assert tuple(header_texts) == tokens.PRICING_HEADERS
    assert trs[0].find(qn("w:trPr")).find(qn("w:tblHeader")) is not None   # header repeats
    for tr in trs:
        assert tr.find(qn("w:trPr")).find(qn("w:cantSplit")) is not None   # rows do not split
    total_cells = trs[-1].findall(qn("w:tc"))
    assert len(total_cells) == 2 and total_cells[1].find(qn("w:tcPr")).find(qn("w:gridSpan")).get(qn("w:val")) == "4"
    xml = tbl.xml if hasattr(tbl, "xml") else ""
    from lxml import etree
    xml = etree.tostring(tbl).decode().lower()
    assert tokens.BRAND_PURPLE.lower() in xml and tokens.FONT.lower() in xml
    assert "calibri" not in xml                                            # no legacy font


def test_generated_proposal_has_token_styled_pricing_table_and_no_lint_errors(tmp_path):
    out = tmp_path / "p.docx"
    gen.generate_proposal(company_name="Acme", proposal_ref="EN-T", logo_path=None, output_path=out,
                          annexure_order=["web_app_pt", "cloud_pt"])
    doc = Document(str(out))
    table = gen._find_assessment_pricing_table(doc.element)
    header = ["".join(t.text for t in tc.iter(qn("w:t"))) for tc in table.findall(qn("w:tr"))[0].findall(qn("w:tc"))]
    assert header[-1] == "Cost Per Annexure"
    assert len(table.findall(qn("w:tr"))) == 4                             # header + 2 + total
    assert lint.lint_docx(doc).ok


# ----------------------------------------------------------------------------- page-break fixes
def _doc_with_body(*paragraph_specs):
    doc = Document()
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    sect = body.find(qn("w:sectPr"))
    for spec in paragraph_specs:
        el = gen._make_page_break_paragraph() if spec == "BR" else (OxmlElement("w:p") if spec == "" else _para(spec))
        sect.addprevious(el)
    return doc


def test_page_break_before_val_zero_is_off():
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    node = OxmlElement("w:pageBreakBefore")
    node.set(qn("w:val"), "0")
    ppr.append(node)
    assert gen._has_page_break_before(ppr) is False
    node.set(qn("w:val"), "1")
    assert gen._has_page_break_before(ppr) is True
    node.attrib.pop(qn("w:val"))
    assert gen._has_page_break_before(ppr) is True
    assert gen._has_page_break_before(None) is False


def test_stacked_page_breaks_are_removed_even_with_blank_paragraphs_between():
    doc = _doc_with_body("Intro", "BR", "", "BR", "Annexure 1: X", "BR", "Annexure 2: Y")
    assert gen._remove_stacked_page_breaks(doc) == 1
    assert [p.text for p in doc.paragraphs if p.text] == ["Intro", "Annexure 1: X", "Annexure 2: Y"]
    assert lint.check_blank_spacing(doc) == []


def test_lint_flags_stacked_breaks_trailing_break_and_blank_runs():
    doc = _doc_with_body("Intro", "BR", "BR", "Body", "", "", "", "Tail", "BR")
    codes = {f.code for f in lint.check_blank_spacing(doc)}
    assert codes == {"stacked_page_breaks", "blank_paragraph_run", "trailing_page_break"}


def test_lint_flags_wrong_annexure_numbering():
    ok = _doc_with_body("Annexure 1: A", "Annexure 2: B")
    bad = _doc_with_body("Annexure 1: A", "Annexure 3: B", "Annexure 3: C")
    assert lint.check_annexure_numbering(ok) == []
    assert [f.code for f in lint.check_annexure_numbering(bad)] == ["annexure_numbering"]


def test_lint_flags_tables_wider_than_the_text_area():
    doc = _doc_with_body("x")
    wide = pp.build_pricing_table([], 20000)
    doc.element.body.find(qn("w:sectPr")).addprevious(wide)
    assert [f.code for f in lint.check_tables_within_width(doc)] == ["table_too_wide"]
    doc2 = _doc_with_body("x")
    doc2.element.body.find(qn("w:sectPr")).addprevious(pp.build_pricing_table([], 8000))  # 5.56 in < 6.0 in
    assert lint.check_tables_within_width(doc2) == []


def test_lint_reports_fonts_and_colours_outside_the_tokens_as_warnings():
    doc = Document()
    run = doc.add_paragraph().add_run("hello")
    run.font.name = "Calibri"
    from docx.shared import RGBColor
    run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    codes = {f.code for f in lint.check_run_formatting(doc)}
    assert codes == {"font_outside_tokens", "color_outside_tokens"}


def test_tokens_are_complete_and_consistent():
    assert tokens.rgb(tokens.BRAND_PURPLE) == (0x70, 0x30, 0xA0)
    assert tokens.BRAND_PURPLE in tokens.ALLOWED_COLORS and tokens.BODY_TEXT in tokens.ALLOWED_COLORS
    assert sum(tokens.PRICING_COLUMN_WIDTH_PCT) == 100
    assert len(tokens.PRICING_HEADERS) == len(tokens.PRICING_ALIGN) == len(tokens.PRICING_COLUMN_WIDTH_PCT) == 5
    assert tokens.COVER_NAVY is None          # unconfirmed brand value is never guessed
