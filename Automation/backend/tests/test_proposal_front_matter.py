"""
Tests for the organisation's fixed front pages (the master "CyberSmithSECURE Annexures 2.docx"):
cover with date and client name, Contents as a real TOC field written from the headings, the eight
fixed pages kept as designed, the master's header/footer, the logo slot, and the proposal date.

Run from Automation/backend/:   python -m pytest tests/test_proposal_front_matter.py -v
"""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

import css_proposal_generator as gen  # noqa: E402
import proposal_contents as contents  # noqa: E402
import proposal_lint as lint  # noqa: E402
import proposal_workflow as wf  # noqa: E402

FIXED_HEADINGS = ["INTRODUCTION", "OUR UNDERSTANDING OF YOUR REQUIREMENT", "TEAM OVERVIEW", "METHODOLOGY OVERVIEW",
                  "VULNERABILITY ASSESSMENT METHODOLOGY"]


def _text(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


@pytest.fixture(scope="module")
def generated(tmp_path_factory) -> Path:
    out = tmp_path_factory.mktemp("front") / "p.docx"
    gen.generate_proposal(company_name="Meridian Textiles Pvt. Ltd.", proposal_ref="CSS-ENG-2026-227",
                          proposal_date="25-10-2026", logo_path=None, output_path=out,
                          annexure_order=["web_app_pt", "mobile_app_pt", "post_breach"])
    return out


def _front(doc) -> list:
    """Top-level body elements from INTRODUCTION to the first annexure title (exclusive)."""
    out, on = [], False
    for el in doc.element.body:
        t = " ".join(_text(el).split()) if el.tag == qn("w:p") else ""
        if t == "INTRODUCTION":
            on = True
        elif on and re.match(r"^Annexure \d+:", t):
            break
        if on:
            out.append(el)
    return out


# ----------------------------------------------------------------------------- cover
def test_cover_shows_the_date_and_the_client_name_not_the_masters_placeholders(generated):
    with zipfile.ZipFile(generated) as z:
        xml = z.read("word/document.xml").decode("utf8")
    text = re.sub(r"<[^>]+>", "", xml)
    assert "25-10-2026" in text and "Meridian Textiles Pvt. Ltd." in text
    assert "15-09-2025" not in text and "Company name" not in text and "{{" not in text


def test_reference_number_is_kept_off_the_cover(generated):
    with zipfile.ZipFile(generated) as z:
        assert "CSS-ENG-2026-227" not in re.sub(r"<[^>]+>", "", z.read("word/document.xml").decode("utf8"))


def test_the_cover_ends_with_a_real_page_break_not_blank_filler(generated):
    doc = Document(str(generated))
    children = list(doc.element.body)
    for_idx = next(i for i, el in enumerate(children) if el.tag == qn("w:p") and _text(el).strip() == "For")
    assert for_idx < 6                                           # the cover paragraph, then "For" (not ~50 empty paragraphs)
    assert children[for_idx].find(f"{qn('w:pPr')}/{qn('w:pageBreakBefore')}") is not None


# ----------------------------------------------------------------------------- header and footer
def test_every_page_carries_the_masters_header_and_footer(generated):
    with zipfile.ZipFile(generated) as z:
        raw = {n: z.read(n).decode("utf8") for n in z.namelist() if re.match(r"word/(header|footer)\d*\.xml$", n)}
    headers = {n: x for n, x in raw.items() if "header" in n}
    footers = {n: x for n, x in raw.items() if "footer" in n}
    assert headers and footers
    assert any("Statement of Work" in re.sub(r"<[^>]+>", " ", x) for x in headers.values())
    assert any("<w:drawing" in x or "<w:pict" in x for x in footers.values())       # the contact strip + purple band
    assert len(Document(str(generated)).sections) == 1           # one section: nothing can lose them mid-document


# ----------------------------------------------------------------------------- fixed front pages
def test_fixed_front_pages_keep_the_masters_sections_in_order(generated):
    doc = Document(str(generated))
    headings = [" ".join(_text(el).split()) for el in _front(doc)
                if el.tag == qn("w:p") and " ".join(_text(el).split()) in FIXED_HEADINGS]
    assert headings == FIXED_HEADINGS
    texts = " ".join(_text(el) for el in _front(doc))
    for phrase in ("CERT-In empanelled", "OSSTMM", "PLANNING PHASE", "DISCOVERY", "Typical tools that may be used"):
        assert phrase in texts


def test_fixed_front_pages_keep_their_diagrams_and_team_table(generated):
    front = _front(Document(str(generated)))
    drawings = sum(len(list(el.iter(qn("w:drawing")))) for el in front)
    assert drawings >= 10                                        # service columns, workflow, frameworks, tools ...
    tables = [el for el in front if el.tag == qn("w:tbl")]
    assert len(tables) == 1 and "Sr No." in _text(tables[0])


def test_each_fixed_page_starts_a_new_page_without_relying_on_filler(generated):
    doc = Document(str(generated))
    starts = {"OUR UNDERSTANDING OF YOUR REQUIREMENT", "METHODOLOGY OVERVIEW", "VULNERABILITY ASSESSMENT METHODOLOGY",
              "PLANNING PHASE", "The vulnerability assessment/penetration testing report consist of:"}
    found = set()
    for el in _front(doc):
        t = " ".join(_text(el).split())
        if el.tag == qn("w:p") and t in starts:
            assert gen._has_page_break_before(el.find(qn("w:pPr"))), t
            found.add(t)
    assert found == starts


def test_front_pages_are_exempt_from_the_blank_line_lint_but_the_rest_is_clean(generated):
    report = lint.lint_docx(generated)
    assert report.ok and report.warnings == [], [str(f) for f in report.findings]


# ----------------------------------------------------------------------------- contents
def test_contents_is_a_real_toc_field_listing_exactly_the_body_headings(generated):
    doc = Document(str(generated))
    entries = lint._contents_entries(doc)
    assert entries == [
        *FIXED_HEADINGS,
        "Annexure 1: Web Application Penetration Testing Assessment",
        "Annexure 2: Mobile Application Penetration Testing",
        "Annexure 3: Post Breach Scenario Assessment",
        "Terms & Conditions", "Conclusion & Pricing"]
    block = contents.contents_block(doc.element.body)
    kinds = [c.get(qn("w:fldCharType")) for c in block[0].iter(qn("w:fldChar"))]
    assert kinds[:2] == ["begin", "separate"]                    # the TOC field opens on the first entry
    assert [c.get(qn("w:fldCharType")) for c in block[-1].iter(qn("w:fldChar"))] == ["end"]
    instr = "".join(i.text or "" for i in block[0].iter(qn("w:instrText")))
    assert instr.strip().startswith("TOC")
    # every entry links to a bookmark on its heading and carries a PAGEREF to it
    bookmarks = {b.get(qn("w:name")) for b in doc.element.body.iter(qn("w:bookmarkStart"))}
    for p in block[:-1]:
        anchor = p.find(qn("w:hyperlink")).get(qn("w:anchor"))
        assert anchor in bookmarks
        assert f"PAGEREF {anchor}" in "".join(i.text or "" for i in p.iter(qn("w:instrText")))


def test_annexure_entries_use_the_annexure_toc_level_and_the_rest_the_first(generated):
    doc = Document(str(generated))
    styles = {}
    for p in contents.contents_block(doc.element.body)[:-1]:
        label = next(iter(p.find(qn("w:hyperlink")).iter(qn("w:t")))).text
        styles[label] = p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}").get(qn("w:val"))
    assert styles["INTRODUCTION"] == "TOC1" and styles["Conclusion & Pricing"] == "TOC1"
    assert styles["Annexure 1: Web Application Penetration Testing Assessment"] == "TOC2"


def test_rebuilding_the_contents_is_idempotent_and_keeps_cached_pages(generated):
    doc = Document(str(generated))
    before = contents.read_pages(doc)
    contents.build(doc)
    assert contents.read_pages(doc) == before
    assert lint._contents_entries(doc).count("INTRODUCTION") == 1


def test_contents_lists_the_extra_annexures_in_template_order(tmp_path):
    out = tmp_path / "all.docx"
    gen.generate_proposal(company_name="All Co", proposal_ref="EN-1", logo_path=None, output_path=out)
    entries = lint._contents_entries(Document(str(out)))
    annexures = [e for e in entries if e.startswith("Annexure ")]
    assert len(annexures) == 13
    assert annexures[10].endswith("Source Code Review") and annexures[11].endswith("Post Breach Scenario Assessment")
    assert annexures[12].endswith("Active Directory Assessment")
    assert lint.lint_docx(out).ok


# ----------------------------------------------------------------------------- logo slot
def test_client_logo_goes_between_for_and_the_client_name(tmp_path):
    from PIL import Image
    logo = tmp_path / "logo.png"
    Image.new("RGB", (240, 90), color=(30, 144, 255)).save(logo)
    out = tmp_path / "with_logo.docx"
    gen.generate_proposal(company_name="Logo Co", proposal_ref="EN-1", logo_path=str(logo), output_path=out,
                          annexure_order=["web_app_pt"])
    doc = Document(str(out))
    paragraphs = doc.paragraphs
    for_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "For")
    name_idx = next(i for i in range(for_idx + 1, len(paragraphs)) if "Logo Co" in paragraphs[i].text)
    between = paragraphs[for_idx + 1:name_idx]
    assert any(p._element.find(f".//{qn('w:drawing')}") is not None for p in between)


def test_without_a_client_logo_the_slot_is_empty_not_a_cybersmith_logo(generated):
    doc = Document(str(generated))
    paragraphs = doc.paragraphs
    for_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "For")
    name_idx = next(i for i in range(for_idx + 1, len(paragraphs)) if "Meridian" in paragraphs[i].text)
    assert not any(p._element.find(f".//{qn('w:drawing')}") is not None for p in paragraphs[for_idx + 1:name_idx])
    assert "None" not in paragraphs[for_idx + 1].text


# ----------------------------------------------------------------------------- source template and catalog
def test_catalog_matches_the_source_template_annexure_by_annexure():
    doc = Document(str(gen.SOURCE_TEMPLATE))
    headings = {}
    for p in doc.paragraphs:
        m = re.match(r"^Annexure (\d+):\s*(.+?)\s*$", p.text.strip())
        if m and "\t" not in p.text:
            headings[int(m.group(1))] = m.group(2)
    assert sorted(headings) == list(range(1, 14))
    for item in gen.ASSESSMENT_CATALOG:
        assert headings[item["annexure_number"]] == item["title"], item["key"]


# ----------------------------------------------------------------------------- proposal date in the Excel flow
def test_the_draft_carries_an_editable_proposal_date_that_defaults_to_today(fixtures):
    import proposal_excel
    result = proposal_excel.parse_scoping_workbook(fixtures / "samples" / "02_known_only.xlsx")
    draft = wf.build_draft(result, "02.xlsx", gen.CATALOG_BY_KEY)
    assert re.fullmatch(r"\d{2}-\d{2}-\d{4}", draft["client"]["proposal_date"])
    edited = wf.merge_edits(draft, {"client": {"proposal_date": "01-02-2027"}})
    assert edited["client"]["proposal_date"] == "01-02-2027"
    assert wf.resolve_plan(edited)["proposal_date"] == "01-02-2027"
    blank = wf.merge_edits(draft, {"client": {"proposal_date": " "}})
    assert any("Proposal date is required" in i for i in wf.readiness_issues(blank))
    bad = wf.merge_edits(draft, {"client": {"proposal_date": "<b>1</b>"}})
    assert any("Proposal date contains unsupported" in i for i in wf.readiness_issues(bad))
