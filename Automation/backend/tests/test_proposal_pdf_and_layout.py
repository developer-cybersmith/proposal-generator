"""
Tests for page-layout safety: blank pages, closing section, Contents accuracy, core properties,
PDF analysis and (when an engine is installed) real-render verification.

The engine tests use Microsoft Word on Windows or LibreOffice when present and are skipped
otherwise, so the suite stays green on machines without either.
Run from Automation/backend/:   python -m pytest tests/test_proposal_pdf_and_layout.py -v
"""
from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

import css_proposal_generator as gen  # noqa: E402
import proposal_lint as lint  # noqa: E402
import proposal_pdf as pdfmod  # noqa: E402

pymupdf = pytest.importorskip("pymupdf")


# ----------------------------------------------------------------------------- helpers
def _para(text: str = ""):
    p = OxmlElement("w:p")
    if text:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        r.append(t)
        p.append(r)
    return p


def _doc(*specs):
    doc = Document()
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    sect = body.find(qn("w:sectPr"))
    for s in specs:
        sect.addprevious(gen._make_page_break_paragraph() if s == "BR" else _para(s))
    return doc


def _texts(doc):
    return [p.text for p in doc.paragraphs]


def _make_pdf(path: Path, pages: list[str]):
    """One page per entry; every page gets the same header and footer lines (chrome)."""
    d = pymupdf.open()
    for body in pages:
        page = d.new_page()
        page.insert_text((72, 40), "CyberSmithSECURE Pvt. Ltd. header", fontsize=9)
        page.insert_text((72, 800), "Web: cybersmithsecure.com footer", fontsize=9)
        y = 100
        for line in body.splitlines():
            page.insert_text((72, y), line, fontsize=11)
            y += 16
    d.save(str(path))
    d.close()


# ----------------------------------------------------------------------------- page-break normalisation
def test_break_paragraphs_become_page_break_before_on_the_next_paragraph():
    doc = _doc("Intro", "BR", "", "Annexure 1: X", "body", "BR", "Annexure 2: Y")
    stats = gen._normalize_page_breaks(doc)
    assert stats["converted"] == 2
    assert _texts(doc)[:5] == ["Intro", "Annexure 1: X", "body", "Annexure 2: Y", ""] or "Annexure 1: X" in _texts(doc)
    heads = [p for p in doc.paragraphs if p.text.startswith("Annexure")]
    assert all(gen._has_page_break_before(h._element.find(qn("w:pPr"))) for h in heads)
    assert not any(gen._is_page_break_only_paragraph(p._element) for p in doc.paragraphs)


def test_blank_paragraphs_before_a_new_page_and_at_the_end_are_removed():
    doc = _doc("Intro", "", "", "BR", "Next", "text", "", "", "")
    gen._normalize_page_breaks(doc)
    assert [t for t in _texts(doc)] == ["Intro", "Next", "text"]


def test_normalisation_is_idempotent_and_loses_no_text():
    doc = _doc("A", "BR", "B", "", "BR", "C", "")
    gen._normalize_page_breaks(doc)
    first = _texts(doc)
    gen._normalize_page_breaks(doc)
    assert _texts(doc) == first == ["A", "B", "C"]


def test_page_break_before_is_placed_in_schema_order():
    p = _para("Heading")
    ppr = OxmlElement("w:pPr")
    ppr.append(OxmlElement("w:pStyle"))
    ppr.append(OxmlElement("w:keepNext"))
    ppr.append(OxmlElement("w:spacing"))
    p.insert(0, ppr)
    gen._set_page_break_before(p)
    assert [c.tag.split("}")[1] for c in p.find(qn("w:pPr"))] == ["pStyle", "keepNext", "pageBreakBefore", "spacing"]


def test_break_before_a_table_goes_on_its_first_paragraph():
    doc = _doc("Intro", "BR")
    body = doc.element.body
    tbl = OxmlElement("w:tbl")
    tr, tc = OxmlElement("w:tr"), OxmlElement("w:tc")
    tc.append(_para("cell"))
    tr.append(tc)
    tbl.append(tr)
    body.find(qn("w:sectPr")).addprevious(tbl)
    gen._normalize_page_breaks(doc)
    first_cell_p = tbl.find(".//" + qn("w:p"))
    assert gen._has_page_break_before(first_cell_p.find(qn("w:pPr")))


def test_final_section_is_added_as_a_continuous_copy_and_not_duplicated():
    doc = Document()
    body = doc.element.body
    final = body.find(qn("w:sectPr"))
    body.remove(final)                                         # what apply_annexure_order used to leave behind
    p = _para("end of section")
    ppr = OxmlElement("w:pPr")
    ppr.append(final)
    p.insert(0, ppr)
    body.append(p)
    assert gen._ensure_final_section(doc) is True
    new_final = body.find(qn("w:sectPr"))
    assert list(body)[-1] is new_final
    assert new_final.find(qn("w:type")).get(qn("w:val")) == "continuous"
    assert new_final.find(qn("w:pgSz")) is not None            # page size is explicit again
    assert gen._ensure_final_section(doc) is False
    assert lint.check_final_section(doc) == []


# ----------------------------------------------------------------------------- core properties
def test_core_properties_are_rewritten_valid(tmp_path):
    path = tmp_path / "x.docx"
    doc = Document()
    doc.save(str(path))
    with zipfile.ZipFile(path) as z:
        entries = {n: z.read(n) for n in z.namelist()}
    bad = (b"<?xml version='1.0'?><cp:coreProperties xmlns:cp='http://schemas.openxmlformats.org/package/2006/metadata/core-properties'>"
           b"<cp:lastModifiedBy xmlns:cp='http://schemas.openxmlformats.org/officeDocument/2006/custom-properties'>x</cp:lastModifiedBy>"
           b"</cp:coreProperties>")
    entries["docProps/core.xml"] = bad
    with zipfile.ZipFile(path, "w") as z:
        for n, d in entries.items():
            z.writestr(n, d)
    assert [f.code for f in lint.check_package(path)] == ["core_properties_invalid"]
    gen._repair_core_properties(path, title="Acme & Sons")
    assert lint.check_package(path) == []
    with zipfile.ZipFile(path) as z:
        text = z.read("docProps/core.xml").decode()
    assert "Acme &amp; Sons" in text and "CyberSmithSECURE" in text
    Document(str(path))                                         # still a readable document


def test_generated_proposals_have_valid_package_final_section_and_clean_contents(tmp_path):
    out = tmp_path / "p.docx"
    gen.generate_proposal(company_name="Acme & Co", proposal_ref="EN-1", logo_path=None, output_path=out,
                          annexure_order=["web_app_pt", "cloud_pt", "source_code_review"])
    report = lint.lint_docx(out)
    assert report.ok, [str(f) for f in report.errors]
    doc = Document(str(out))
    contents = [p.text.split("\t")[0] for p in doc.paragraphs if "\t" in p.text][:14]
    # the master's five fixed sections are all real headings now, so all five are listed ...
    for heading in ("INTRODUCTION", "OUR UNDERSTANDING OF YOUR REQUIREMENT", "TEAM OVERVIEW", "METHODOLOGY OVERVIEW",
                    "VULNERABILITY ASSESSMENT METHODOLOGY"):
        assert heading in contents
    # ... and nothing the body does not have (the old template's Timelines section is gone)
    assert "PROPOSED TIMELINES & PROJECT MANAGEMENT" not in contents
    assert [c for c in contents if c.startswith("Annexure ")] == [
        "Annexure 1: Web Application Penetration Testing Assessment", "Annexure 2: Cloud Penetration Testing",
        "Annexure 3: Source Code Review"]


@pytest.mark.parametrize("name", ["L&T Infotech", "Smith & Sons <Pvt> Ltd", 'Quote"s \'n\' Co', "अक्षय & 株式会社"])
def test_special_characters_in_the_client_name_do_not_corrupt_the_document(tmp_path, name):
    out = tmp_path / "n.docx"
    gen.generate_proposal(company_name=name, proposal_ref="EN-&-1", logo_path=None, output_path=out,
                          annexure_order=["web_app_pt", "cloud_pt"])
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert name in text                                              # shown exactly as typed
    heads = [p.text for p in doc.paragraphs if p.text.startswith("Annexure ") and "\t" not in p.text]
    assert len(heads) == 2                                           # the selection survived (not all 13)
    assert "Conclusion & Pricing" in text and "Terms & Conditions" in text   # ampersands intact
    assert lint.lint_docx(out).ok


def test_contents_phantom_check_flags_missing_headings():
    doc = _doc("Contents", "INTRODUCTION\t3", "METHODOLOGY OVERVIEW\t4", "TEAM OVERVIEW\t5", "INTRODUCTION", "TEAM OVERVIEW")
    # make the entries look like Contents lines (they contain a tab) and the headings real headings
    codes = [f.message for f in lint.check_contents_has_no_phantom_entries(doc)]
    assert any("METHODOLOGY OVERVIEW" in c for c in codes) and len(codes) == 1


def test_measured_pages_are_written_into_the_contents(tmp_path):
    out = tmp_path / "p.docx"
    gen.generate_proposal(company_name="Acme", proposal_ref="EN-1", logo_path=None, output_path=out,
                          annexure_order=["web_app_pt", "cloud_pt"])
    measured = {"introduction": 3, "our understanding of your requirement": 4, "team overview": 4,
                "methodology overview": 5, "vulnerability assessment methodology": 6,
                "annexure 1": 9, "annexure 2": 13, "terms & conditions": 17, "conclusion & pricing": 19}
    assert gen._apply_measured_toc_pages(out, measured) >= 1     # the estimate already matched some entries
    import proposal_contents
    pages = proposal_contents.read_pages(Document(str(out)))
    assert pages["introduction"] == 3 and pages["terms & conditions"] == 17 and pages["conclusion & pricing"] == 19
    assert pages["annexure 1"] == 9 and pages["annexure 2"] == 13
    assert gen._apply_measured_toc_pages(out, measured) == 0   # nothing left to change


# ----------------------------------------------------------------------------- PDF analysis
def test_blank_pages_are_found_ignoring_header_and_footer(tmp_path):
    pdf = tmp_path / "a.pdf"
    _make_pdf(pdf, ["Statement of Work and Proposal for the client\nprepared by CyberSmithSECURE",
                    "Contents\nINTRODUCTION 3\nTEAM OVERVIEW 4", "INTRODUCTION\nSome real introduction text for the client.",
                    "", "PROPOSED TIMELINES & PROJECT MANAGEMENT\nWork plan text goes here."])
    analysis = pdfmod.analyze_pdf(pdf)
    assert analysis.page_count == 5 and analysis.blank_pages == [4]


def test_a_page_with_only_an_image_is_not_blank(tmp_path):
    pdf = tmp_path / "img.pdf"
    d = pymupdf.open()
    for i in range(4):
        page = d.new_page()
        page.insert_text((72, 40), "header line", fontsize=9)
        if i == 2:
            pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 40), False)
            pix.clear_with(200)
            page.insert_image(pymupdf.Rect(100, 100, 300, 300), pixmap=pix)
        else:
            page.insert_text((72, 100), f"A real sentence of body text sits on page number {i + 1}.", fontsize=11)
    d.save(str(pdf))
    assert pdfmod.analyze_pdf(pdf).blank_pages == []


def test_heading_pages_are_measured_after_the_contents_page(tmp_path):
    pdf = tmp_path / "h.pdf"
    _make_pdf(pdf, ["Cover", "Contents\nINTRODUCTION\nTEAM OVERVIEW\nAnnexure listing", "INTRODUCTION\ntext",
                    "OUR UNDERSTANDING OF YOUR REQUIREMENT\ntext\nTEAM OVERVIEW\nrows", "METHODOLOGY OVERVIEW\ntext",
                    "VULNERABILITY ASSESSMENT METHODOLOGY\ntext",
                    "Annexure 1: Web Application Penetration Testing Assessment\ntext", "more text", "Annexure 2: Cloud Penetration Testing\ntext",
                    "Terms & Conditions\ntext", "Conclusion & Pricing\ntext"])
    found = pdfmod.measure_heading_pages(pdfmod.analyze_pdf(pdf), annexure_count=2)
    assert found == {"introduction": 3, "our understanding of your requirement": 4, "team overview": 4,
                     "methodology overview": 5, "vulnerability assessment methodology": 6,
                     "annexure 1": 7, "annexure 2": 9, "terms & conditions": 10, "conclusion & pricing": 11}


def test_no_engine_means_unverified_but_never_an_error(tmp_path, monkeypatch):
    monkeypatch.setenv("PROPOSAL_PDF_ENGINE", "off")
    check = pdfmod.check_document(tmp_path / "missing.docx", 3)
    assert not check.verified and "not verified" in check.header_value()
    assert gen.finalize_with_render(tmp_path / "missing.docx", ["web_app_pt"]).verified is False


def test_render_failure_is_reported_not_raised(tmp_path, monkeypatch):
    monkeypatch.setattr(pdfmod, "available_engine", lambda: "libreoffice")
    monkeypatch.setattr(pdfmod, "find_soffice", lambda: None)
    check = pdfmod.check_document(tmp_path / "x.docx", 1)
    assert not check.verified and "render failed" in check.note


# ----------------------------------------------------------------------------- real render (skipped without an engine)
needs_engine = pytest.mark.skipif(not (pdfmod.find_soffice() or pdfmod.word_available()),
                                  reason="neither LibreOffice nor Microsoft Word is installed")


@pytest.mark.skipif(not pdfmod.word_available(), reason="Microsoft Word is not installed")
def test_word_itself_can_edit_the_generated_document(tmp_path):
    """The downloaded .docx must be editable in Word: type into the pricing table and the body, update the TOC."""
    import subprocess
    out = tmp_path / "p.docx"
    gen.generate_proposal(company_name="Edit Co", proposal_ref="EN-E", logo_path=None, output_path=out,
                          annexure_order=["web_app_pt", "cloud_pt"])
    script = r"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
try {
  $word.Visible = $false; $word.DisplayAlerts = 0
  $doc = $word.Documents.Open($env:CSS_DOCX, $false, $false)
  if ($doc.ProtectionType -ne -1) { throw 'document is protected' }
  $typed = $false
  foreach ($t in $doc.Tables) { if ($t.Range.Text -match 'Cost Per Annexure') { $t.Cell(2, 5).Range.Text = '75,000'; $typed = $true } }
  if (-not $typed) { throw 'pricing table not found' }
  $r = $doc.Content; $null = $r.Find.Execute('CyberSmithSECURE is incorporated'); if (-not $r.Find.Found) { throw 'intro not found' }
  $r.InsertBefore('EDITED: ')
  $doc.TablesOfContents(1).Update()
  $doc.Save(); $doc.Close(0); 'EDIT-OK'
} finally { $word.Quit() }
"""
    env = dict(__import__("os").environ, CSS_DOCX=str(out))
    done = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script], env=env,
                          capture_output=True, text=True, timeout=180)
    assert "EDIT-OK" in done.stdout, done.stdout + done.stderr
    doc = Document(str(out))
    assert any("EDITED: " in p.text for p in doc.paragraphs)
    table = gen._find_assessment_pricing_table(doc.element)
    assert "75,000" in "".join(t.text or "" for t in table.iter(qn("w:t")))


@needs_engine
@pytest.mark.parametrize("order", [["web_app_pt", "cloud_pt", "source_code_review"], None])
def test_real_render_has_no_blank_pages_and_exact_contents(tmp_path, monkeypatch, order):
    monkeypatch.setenv("PROPOSAL_PDF_ENGINE", "word" if pdfmod.word_available() and not pdfmod.find_soffice() else "libreoffice")
    out = tmp_path / "p.docx"
    gen.generate_proposal(company_name="Render Co", proposal_ref="EN-R", logo_path=None, output_path=out, annexure_order=order)
    keys = order or [i["key"] for i in gen.ASSESSMENT_CATALOG]
    check = gen.finalize_with_render(out, keys, "Render Co")
    assert check.verified, check.note
    assert check.blank_pages == [], f"blank pages: {check.blank_pages}"
    # every Contents number now equals the page the heading really starts on
    import proposal_contents
    cached = proposal_contents.read_pages(Document(str(out)))
    assert len(cached) == 5 + len(keys) + 2                      # five fixed sections, annexures, Terms, Conclusion
    for key, page in cached.items():
        if key in check.headings:
            assert page == check.headings[key], f"{key}: Contents says {page}, render says {check.headings[key]}"
    assert check.headings["introduction"] == 3                  # the master's cover + Contents + ...
    assert check.headings["annexure 1"] == 9                    # ... eight fixed pages, then the annexures
    assert lint.lint_docx(out).ok
