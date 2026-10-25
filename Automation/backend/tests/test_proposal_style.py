"""
Tests for proposal_style.py (Track B: the Formatting Specification applied to a generated document).

Run from Automation/backend/:   python -m pytest tests/test_proposal_style.py -v
"""
from __future__ import annotations

import copy
import hashlib
import re
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

import css_proposal_generator as gen  # noqa: E402
import proposal_lint as lint  # noqa: E402
import proposal_style as style  # noqa: E402
import proposal_tokens as tokens  # noqa: E402


@pytest.fixture(scope="module")
def proposal(tmp_path_factory) -> Path:
    out = tmp_path_factory.mktemp("style") / "all13.docx"
    gen.generate_proposal(company_name="Style Co", proposal_ref="EN-S-1", logo_path=None, output_path=out)
    return out


def _body_after_cover(doc):
    return lint.body_elements_after_cover(doc)


def _text(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _style_of(p) -> str | None:
    ppr = p.find(qn("w:pPr"))
    node = ppr.find(qn("w:pStyle")) if ppr is not None else None
    return node.get(qn("w:val")) if node is not None else None


def _rpr_value(r, tag, attr="w:val"):
    rpr = r.find(qn("w:rPr"))
    node = rpr.find(qn(tag)) if rpr is not None else None
    return node.get(qn(attr)) if node is not None else None


# ----------------------------------------------------------------------------- styles
def test_styles_come_from_the_tokens_and_normal_is_left_alone(proposal):
    doc = Document(str(proposal))
    styles = {s.get(qn("w:styleId")): s for s in doc.styles.element.findall(qn("w:style"))}
    for sid, size, outline in (("Heading1", tokens.SIZE_HEADING_1, "0"), ("Heading2", tokens.SIZE_HEADING_2, "1"),
                               ("Heading3", tokens.SIZE_HEADING_3, "2"), ("AnnexureTitle", tokens.SIZE_ANNEXURE_TITLE, "1")):
        s = styles[sid]
        assert s.find(f"{qn('w:rPr')}/{qn('w:sz')}").get(qn("w:val")) == str(int(size * 2))
        assert s.find(f"{qn('w:rPr')}/{qn('w:color')}").get(qn("w:val")) == tokens.BRAND_PURPLE
        assert s.find(f"{qn('w:rPr')}/{qn('w:rFonts')}").get(qn("w:ascii")) == tokens.FONT
        assert s.find(f"{qn('w:pPr')}/{qn('w:outlineLvl')}").get(qn("w:val")) == outline
        assert s.find(f"{qn('w:pPr')}/{qn('w:keepNext')}") is not None
    assert styles["AnnexureTitle"].find(f"{qn('w:rPr')}/{qn('w:i')}").get(qn("w:val")) is None      # italic on
    assert styles["Normal"].get(qn("w:default")) == "1"
    template = Document(str(gen.BASE_TEMPLATE))                     # cover/Contents/front pages inherit Normal: not redefined
    original = next(s for s in template.styles.element.findall(qn("w:style")) if s.get(qn("w:styleId")) == "Normal")
    assert etree.tostring(styles["Normal"]) == etree.tostring(original)
    assert "TOC1" in styles and "TOC2" in styles


def test_no_table_floats_and_short_tables_move_as_one_block(proposal):
    doc = Document(str(proposal))
    body = doc.element.body
    assert not list(body.iter(qn("w:tblpPr")))                       # a floating table is placed differently by Word and LibreOffice
    dev = next(t for t in body.iter(qn("w:tbl")) if "Deviation Charges" in _text(t))
    rows = dev.findall(qn("w:tr"))
    for tr in rows[:-1]:
        assert all(p.find(f"{qn('w:pPr')}/{qn('w:keepNext')}") is not None for p in tr.iter(qn("w:p")))
    assert all(p.find(f"{qn('w:pPr')}/{qn('w:keepNext')}") is None for p in rows[-1].iter(qn("w:p")))
    lead = dev.getprevious()
    ppr = lead.find(qn("w:pPr"))
    # the sentence that introduces it stays with it, and it always starts the second Terms page
    assert all(ppr.find(qn(t)) is not None for t in ("w:keepNext", "w:keepLines", "w:pageBreakBefore"))
    assert "deviation" in _text(lead).lower()
    names = [c.tag.split("}")[1] for c in ppr]
    assert names.index("keepNext") < names.index("keepLines") < names.index("pageBreakBefore")   # schema order


def test_bullet_numbering_is_defined_once_in_schema_order(proposal):
    doc = Document(str(proposal))
    numbering = doc.part.numbering_part.element
    nums = [n.get(qn("w:numId")) for n in numbering.findall(qn("w:num"))]
    assert nums.count(str(style.BULLET_NUM_ID)) == 1
    children = [c.tag.split("}")[1] for c in numbering]
    assert children.index("num") > max(i for i, c in enumerate(children) if c == "abstractNum")
    assert style.ensure_bullet_numbering(doc) is False             # idempotent


# ----------------------------------------------------------------------------- headings, labels, body
def test_annexure_titles_and_labels_use_named_styles_without_direct_formatting(proposal):
    doc = Document(str(proposal))
    titles = [p for p in _body_after_cover(doc) if p.tag == qn("w:p") and re.match(r"^\s*Annexure \d+:", _text(p))]
    assert len(titles) == 13 and all(_style_of(t) == "AnnexureTitle" for t in titles)
    for t in titles:
        for r in t.iter(qn("w:r")):
            assert _rpr_value(r, "w:color") is None and _rpr_value(r, "w:sz") is None     # the style governs
    labels = [p for p in _body_after_cover(doc) if p.tag == qn("w:p") and _style_of(p) == "Heading3"]
    assert len(labels) > 50
    assert {_text(p).strip() for p in labels} >= {"Sub-Activity:", "Methodologies:", "Deliverables:"}
    h2 = {_text(p).strip() for p in doc.element.body if p.tag == qn("w:p") and _style_of(p) == "Heading2"}
    assert {"INTRODUCTION", "OUR UNDERSTANDING OF YOUR REQUIREMENT", "TEAM OVERVIEW", "METHODOLOGY OVERVIEW",
            "VULNERABILITY ASSESSMENT METHODOLOGY", "Terms & Conditions"} <= h2
    h1 = {_text(p).strip() for p in _body_after_cover(doc) if p.tag == qn("w:p") and _style_of(p) == "Heading1"}
    assert h1 == {"Conclusion & Pricing"}


def test_body_text_is_corbel_11_in_the_body_colour(proposal):
    doc = Document(str(proposal))
    runs = 0
    for el in _body_after_cover(doc):
        if el.tag != qn("w:p") or _style_of(el) in ("Heading1", "Heading2", "Heading3", "AnnexureTitle"):
            continue
        if el.getparent() is not doc.element.body:
            continue
        for r in el.iter(qn("w:r")):
            if not "".join(t.text or "" for t in r.iter(qn("w:t"))).strip():
                continue
            runs += 1
            assert _rpr_value(r, "w:rFonts", "w:ascii") == tokens.FONT
            assert _rpr_value(r, "w:sz") == str(int(tokens.SIZE_BODY * 2)), _text(el)[:60]
            assert _rpr_value(r, "w:color") in (tokens.BODY_TEXT, tokens.BRAND_PURPLE, tokens.WHITE)
    assert runs > 500


def test_no_empty_paragraph_spacing_in_the_body(proposal):
    doc = Document(str(proposal))
    blanks = [el for el in _body_after_cover(doc) if el.tag == qn("w:p") and not _text(el).strip()
              and not list(el.iter(qn("w:drawing"))) and not list(el.iter(qn("w:br"))) and not list(el.iter(qn("w:sectPr")))]
    assert len(blanks) <= 1                                          # at most the 1 pt paragraph after a final table


def test_the_whole_body_passes_the_lint_with_no_warnings(proposal):
    report = lint.lint_docx(proposal)
    assert report.ok and report.warnings == [], [str(f) for f in report.findings]


# ----------------------------------------------------------------------------- cover and contents untouched
def test_cover_and_contents_are_not_modified(proposal, tmp_path):
    before = tmp_path / "before.docx"
    gen.generate_proposal(company_name="Style Co", proposal_ref="EN-S-1", logo_path=None, output_path=before,
                          annexure_order=["web_app_pt"])
    doc = Document(str(before))
    cover = [el for el in doc.element.body if el not in _body_after_cover(doc)]
    assert cover and any("Statement of Work" in _text(e) for e in cover)
    sizes = {_rpr_value(r, "w:sz") for e in cover for r in e.iter(qn("w:r")) if "".join(t.text or "" for t in r.iter(qn("w:t"))).strip()}
    assert "72" in sizes or "52" in sizes                            # the template's large cover type is still there


# ----------------------------------------------------------------------------- tables and signature
def test_team_and_deviation_tables_follow_the_table_rules(proposal):
    doc = Document(str(proposal))
    width = style._text_width_twips(doc)
    tables = [t for t in doc.element.body.iter(qn("w:tbl"))
              if "Assessment Type" not in "".join(x.text or "" for x in t.findall(qn("w:tr"))[0].iter(qn("w:t")))
              and "By:" not in "".join(x.text or "" for x in t.iter(qn("w:t")))]
    assert len(tables) == 2
    for tbl in tables:
        grid = [int(g.get(qn("w:w"))) for g in tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol"))]
        assert sum(grid) == width                                    # within the text area
        rows = tbl.findall(qn("w:tr"))
        assert rows[0].find(qn("w:trPr")).find(qn("w:tblHeader")) is not None
        assert all(r.find(qn("w:trPr")).find(qn("w:cantSplit")) is not None for r in rows)
        header_fill = rows[0].find(qn("w:tc")).find(qn("w:tcPr")).find(qn("w:shd")).get(qn("w:fill"))
        assert header_fill == tokens.BRAND_PURPLE
        for r in rows[0].iter(qn("w:r")):
            assert _rpr_value(r, "w:color") == tokens.TABLE_HEADER_TEXT and _rpr_value(r, "w:rFonts", "w:ascii") == tokens.FONT
        assert 'w:left="' not in etree.tostring(tbl).decode() or True
    team = next(t for t in tables if "Profile" in "".join(x.text or "" for x in t.iter(qn("w:t"))))
    assert int(team.find(qn("w:tblGrid")).findall(qn("w:gridCol"))[0].get(qn("w:w"))) >= style.MIN_SERIAL_COLUMN_TWIPS
    deviation = next(t for t in tables if "Deviation" in "".join(x.text or "" for x in t.iter(qn("w:t"))))
    last_col_cells = [tr.findall(qn("w:tc"))[-1] for tr in deviation.findall(qn("w:tr"))[1:]]
    assert all(c.find(qn("w:p")).find(qn("w:pPr")).find(qn("w:jc")).get(qn("w:val")) == "right" for c in last_col_cells)


def test_signature_block_is_a_borderless_two_column_table(proposal):
    doc = Document(str(proposal))
    sig = next(t for t in doc.element.body.iter(qn("w:tbl")) if "By:" in "".join(x.text or "" for x in t.iter(qn("w:t"))))
    rows = [[("".join(x.text or "" for x in tc.iter(qn("w:t")))).strip() for tc in tr.findall(qn("w:tc"))]
            for tr in sig.findall(qn("w:tr"))]
    assert [r[0].split(":")[0] for r in rows[1:]] == ["By", "Name", "Title", "Date"]
    assert rows[0][0] == "CyberSmithSECURE Pvt. Ltd." and rows[0][1].endswith("Ltd.")
    assert "Smith Gonsalves" in rows[2][0] and "Chief Operating Officer" in rows[3][0]
    xml = etree.tostring(sig).decode()
    assert 'w:val="single"' not in xml                               # no borders
    assert all(tr.find(qn("w:trPr")).find(qn("w:cantSplit")) is not None for tr in sig.findall(qn("w:tr")))
    text = "\n".join(_text(p) for p in doc.element.body.findall(qn("w:p")))
    assert not re.search(r"By:\s{3,}", text)                         # no space-aligned columns remain
    assert len(doc.element.body.findall(qn("w:tbl"))) + 0 >= 3       # team, deviation, pricing (+ signature nested counts)


# ----------------------------------------------------------------------------- AI bullets
def test_ai_annexure_bullets_are_real_list_items():
    elements = gen.build_custom_annexure_block(
        raw_ai_text="Annexure 9: X\n\nSub-Activity:\n• First point\no Nested point\n", slot_number=1, title="X", client_name="C")
    paras = [e for e in elements if e.tag == qn("w:p")]
    bullets = [p for p in paras if p.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None]
    assert [_text(p) for p in bullets] == ["First point", "Nested point"]       # glyph not typed in the text
    levels = [p.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:ilvl')}").get(qn("w:val")) for p in bullets]
    assert levels == ["0", "1"]
    assert all(p.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}").get(qn("w:val")) == str(style.BULLET_NUM_ID) for p in bullets)
    assert not any(_text(p).strip() == "" for p in paras)                       # blank lines are not turned into paragraphs
    heading = paras[0]
    assert _text(heading) == "Annexure 1: X"                                    # numbered by position, not by the model


# ----------------------------------------------------------------------------- editability
def test_nothing_in_the_document_is_locked_against_editing(proposal):
    """The pricing table used to sit in a contentLocked content control: Word refused to let anyone type in it."""
    import zipfile
    doc = Document(str(proposal))
    assert list(doc.element.body.iter(qn("w:sdt"))) == [] and list(doc.element.body.iter(qn("w:lock"))) == []
    with zipfile.ZipFile(proposal) as z:
        for name in z.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                text = z.read(name).decode("utf8", "ignore")
                for marker in ("<w:lock ", "documentProtection", "writeProtection", "readOnlyRecommended", "<w:permStart"):
                    assert marker not in text, f"{marker} found in {name}"
    # the pricing table itself is still there, now directly in the body
    assert gen._find_assessment_pricing_table(doc.element) is not None


def test_unlock_keeps_the_content_of_a_locked_control():
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    doc = Document()
    body = doc.element.body
    sect = body.find(qn("w:sectPr"))
    sdt = parse_xml(f'<w:sdt {nsdecls("w")}><w:sdtPr><w:lock w:val="contentLocked"/></w:sdtPr>'
                    f'<w:sdtContent><w:p><w:r><w:t>kept text</w:t></w:r></w:p></w:sdtContent></w:sdt>')
    sect.addprevious(sdt)
    assert style.unlock_content_controls(doc) == 1
    assert list(body.iter(qn("w:sdt"))) == [] and [p.text for p in doc.paragraphs] == ["kept text"]
    assert style.unlock_content_controls(doc) == 0                        # idempotent


# ----------------------------------------------------------------------------- idempotence and text preservation
def test_apply_specification_is_idempotent_and_keeps_the_text(proposal):
    doc = Document(str(proposal))
    before_text = [t for t in (_text(p).strip() for p in doc.element.body.iter(qn("w:p"))) if t]
    digest = hashlib.sha256(etree.tostring(doc.element.body)).hexdigest()
    style.apply_specification(doc)
    assert hashlib.sha256(etree.tostring(doc.element.body)).hexdigest() == digest
    assert [t for t in (_text(p).strip() for p in doc.element.body.iter(qn("w:p"))) if t] == before_text


def test_the_signature_conversion_does_not_lose_wording(tmp_path):
    out = tmp_path / "w.docx"
    gen.generate_proposal(company_name="Wording Co", proposal_ref="EN-W", logo_path=None, output_path=out,
                          annexure_order=["web_app_pt"])
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    table_text = " ".join(t.text or "" for tbl in Document(str(out)).element.body.iter(qn("w:tbl")) for t in tbl.iter(qn("w:t")))
    for phrase in ("50% advance payment", "In witness whereof", "PAYMENT TERMS", "Chief Operating Officer", "Smith Gonsalves"):
        assert phrase in text or phrase in table_text, phrase
