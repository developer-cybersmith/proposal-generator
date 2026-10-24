"""
Tests for the Excel-driven workflow: drafts, edits, readiness, plan, API, AI retry, generation.

Uses the generated scoping workbooks (see conftest.py) and a fake OpenAI client: no network call
is ever made. Run from Automation/backend/:   python -m pytest tests/test_proposal_workflow.py -v
"""
from __future__ import annotations

import io
import json
import re
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

import css_proposal_generator as gen  # noqa: E402
import proposal_excel  # noqa: E402
import proposal_lint  # noqa: E402
import proposal_workflow as wf  # noqa: E402
from proposal_drafts import DraftStore  # noqa: E402
from proposal_ratelimit import limiter  # noqa: E402

CATALOG = gen.CATALOG_BY_KEY


def sample(fixtures: Path, name: str) -> Path:
    return fixtures / "samples" / f"{name}.xlsx"


def draft_for(fixtures: Path, name: str) -> dict:
    result = proposal_excel.parse_scoping_workbook(sample(fixtures, name))
    assert result.ok, result.errors
    return wf.build_draft(result, f"{name}.xlsx", CATALOG)


def edit(draft: dict, **kw) -> dict:
    return wf.merge_edits(draft, kw)


def entry(draft: dict, entry_id: str) -> dict:
    return next(e for e in draft["entries"] if e["id"] == entry_id)


# ============================================================================ pure functions
def test_draft_for_the_screenshot_example(fixtures):
    d = draft_for(fixtures, "31_image_example_counts")
    web, saas = entry(d, "web_app_vapt"), entry(d, "saas_security")
    assert (web["count"], web["quantity"], web["catalog_key"]) == (6, 23, "web_app_pt")
    assert (saas["count"], saas["quantity"], saas["catalog_key"]) == (3, 12, None)
    assert [e["id"] for e in d["entries"] if e["origin"] == "excel"] == ["web_app_vapt", "saas_security"]
    manual = [e for e in d["entries"] if e["origin"] == "manual"]
    assert {e["annexure"] for e in manual} == {1, 2, 6, 9, 12} and not any(e["include"] for e in manual)
    assert [(r["title"], r["quantity"], r["frequency"], r["cost"]) for r in d["pricing"]] == [
        ("Web Application Penetration Testing Assessment", "23", "One Time", ""),
        ("SaaS Security Assessment", "12", "One Time", "")]


def test_draft_uses_the_default_order_template_annexures_first(fixtures):
    d = draft_for(fixtures, "01_all_in_scope_full")
    included = [e for e in d["entries"] if e["include"]]
    annexures = [e["annexure"] for e in included]
    template = [a for a in annexures if a != "custom"]
    assert template == sorted(template)
    assert annexures[len(template):] == ["custom"] * (len(annexures) - len(template))


def test_merge_edits_whitelists_and_clamps(fixtures):
    d = draft_for(fixtures, "31_image_example_counts")
    out = edit(d, client={"name": "  Acme\x00 Ltd  ", "reference_no": "EN-1", "evil": "x"},
               entries=[{"id": "web_app_vapt", "count": "9999999999", "count_label": "X" * 500, "frequency": "Quarterly\x07"},
                        {"id": "does_not_exist", "include": False}, "garbage", {"id": "saas_security", "title": "Zero Trust Review"}])
    assert out["client"]["name"] == "Acme Ltd" and out["client"]["reference_no"] == "EN-1"
    web = entry(out, "web_app_vapt")
    assert web["count"] == wf.MAX_COUNT and len(web["count_label"]) == 80 and web["frequency"] == "Quarterly"
    assert entry(out, "saas_security")["title"] == "Zero Trust Review"
    assert "evil" not in out["client"] and len(out["entries"]) == len(d["entries"])


def test_template_annexure_title_cannot_be_edited_but_ai_title_can(fixtures):
    d = draft_for(fixtures, "31_image_example_counts")
    out = edit(d, entries=[{"id": "web_app_vapt", "title": "Hacked"}, {"id": "saas_security", "title": "Zero Trust Review"}])
    assert entry(out, "web_app_vapt")["title"] == "Web Application Penetration Testing Assessment"
    assert entry(out, "saas_security")["title"] == "Zero Trust Review"


def test_order_changes_only_when_every_annexure_is_sent(fixtures):
    d = draft_for(fixtures, "31_image_example_counts")
    ids = [e["id"] for e in d["entries"]]
    partial = edit(d, entries=[{"id": "saas_security"}, {"id": "web_app_vapt"}])
    assert [e["id"] for e in partial["entries"]] == ids                       # partial update: unchanged
    full = edit(d, entries=[{"id": i} for i in reversed(ids)])
    assert [e["id"] for e in full["entries"]] == list(reversed(ids))
    assert len(full["entries"]) == len(d["entries"])


def test_external_and_internal_share_one_frequency(fixtures):
    d = draft_for(fixtures, "05_external_and_internal_both")
    out = edit(d, entries=[{"id": "external_network_pt", "frequency": "Quarterly"},
                           {"id": "internal_network_pt", "frequency": "Half-yearly"}])
    assert entry(out, "external_network_pt")["frequency"] == entry(out, "internal_network_pt")["frequency"] == "Quarterly"
    assert all(r["frequency"] == "Quarterly" for r in out["pricing"])           # linked rows follow the annexure


def test_pricing_rows_follow_preview_1_and_keep_user_deletions(fixtures):
    d = draft_for(fixtures, "05_external_and_internal_both")
    assert [r["entry_id"] for r in d["pricing"]] == ["external_network_pt", "internal_network_pt"]   # two lines (D11)
    # the user deletes the Internal row in Preview 2
    out = edit(d, pricing=[r for r in d["pricing"] if r["entry_id"] == "external_network_pt"])
    assert [r["entry_id"] for r in out["pricing"]] == ["external_network_pt"]
    # an unrelated edit keeps it deleted
    out = edit(out, client={"name": "X"})
    assert [r["entry_id"] for r in out["pricing"]] == ["external_network_pt"]
    # excluding then re-including the annexure brings the row back
    out = edit(out, entries=[{"id": "internal_network_pt", "include": False}])
    out = edit(out, entries=[{"id": "internal_network_pt", "include": True}])
    assert [r["entry_id"] for r in out["pricing"]] == ["external_network_pt", "internal_network_pt"]
    # excluding an annexure drops its row
    out = edit(out, entries=[{"id": "external_network_pt", "include": False}])
    assert [r["entry_id"] for r in out["pricing"]] == ["internal_network_pt"]


def test_user_added_pricing_rows_are_kept_and_capped(fixtures):
    d = draft_for(fixtures, "09_minimal")
    rows = list(d["pricing"]) + [{"entry_id": None, "title": "Travel", "quantity": "1", "cost": "5,000"}] * 60
    out = edit(d, pricing=rows)
    assert len(out["pricing"]) <= wf.MAX_ROWS
    assert any(r["title"] == "Travel" and r["entry_id"] is None for r in out["pricing"])


def test_readiness_zero_quantity_needs_confirmation_edit_or_removal(fixtures):
    d = draft_for(fixtures, "10_zero_quantity")
    assert any("nothing above 0" in i.lower() for i in wf.readiness_issues(d))
    for resolve in (
        lambda x: edit(x, entries=[{"id": e["id"], "confirmed": True} for e in x["entries"] if e["include"]]),
        lambda x: edit(x, entries=[{"id": e["id"], "count": 2} for e in x["entries"] if e["include"]]),
        lambda x: edit(x, entries=[{"id": e["id"], "include": False} for e in x["entries"] if e["include"]]
                       + [{"id": "post_breach", "include": True}]),
    ):
        issues = wf.readiness_issues(resolve(d))
        assert not any("nothing above 0" in i.lower() for i in issues), issues


def test_readiness_requires_reference_and_client(fixtures):
    d = draft_for(fixtures, "11_placeholders_untouched")
    issues = wf.readiness_issues(d)
    assert any("Reference number is required" in i for i in issues)
    assert wf.readiness_issues(edit(d, client={"reference_no": "EN-9"})) == []
    assert any("unsupported characters" in i for i in wf.readiness_issues(edit(d, client={"reference_no": "<b>x</b>"})))
    assert any("Client name" in i for i in wf.readiness_issues(edit(d, client={"name": ""})))


def test_readiness_requires_ai_fields_only_when_an_ai_annexure_is_included(fixtures):
    d = draft_for(fixtures, "04_mixed_known_and_custom")
    issues = " ".join(wf.readiness_issues(d))
    assert "Sector is required" in issues
    assert "Duration is required" not in issues                                 # derived from the Excel's start/end dates
    assert d["ai"]["duration"].endswith("days")
    cleared = edit(d, ai={"duration": ""})
    assert "Duration is required" in " ".join(wf.readiness_issues(cleared))
    d = edit(d, ai={"sector": "BFSI"})
    assert not any("AI-written" in i or "Sector" in i for i in wf.readiness_issues(d))
    known_only = draft_for(fixtures, "02_known_only")
    assert not any("Sector" in i for i in wf.readiness_issues(known_only))


def test_sector_is_suggested_from_the_regulatory_driver_but_never_filled(fixtures):
    assert proposal_excel.suggest_sector("RBI") .startswith("Banking")
    assert proposal_excel.suggest_sector("ISO 27001") == ""
    d = draft_for(fixtures, "04_mixed_known_and_custom")
    assert d["ai"]["sector"] == ""                                              # the user confirms (D20)


def test_resolve_plan_merges_external_and_internal_into_one_annexure(fixtures):
    d = draft_for(fixtures, "05_external_and_internal_both")
    plan = wf.resolve_plan(d)
    assert plan["ordered_keys"] == ["network_pt"]
    assert plan["annexure_details"]["network_pt"]["count_text"] == \
        "6 External Infrastructure categories, 12 Internal Infrastructure categories"
    assert [r["title"] for r in plan["pricing_rows"]] == ["External Network Penetration Testing",
                                                          "Internal Network Penetration Testing"]


def test_resolve_plan_gives_ai_annexures_unique_valid_keys(fixtures):
    d = edit(draft_for(fixtures, "03_custom_only"), ai={"sector": "IT", "duration": "10 days"})
    plan = wf.resolve_plan(d)
    assert len(plan["customs"]) == 5 and len(set(plan["ordered_keys"])) == 5
    assert all(gen.OTHER_KEY_RE.match(k) for k in plan["ordered_keys"])


def test_total_cost_for_the_ai_comes_from_the_pricing_table():
    assert wf.total_cost_text([{"cost": "50000"}, {"cost": "1,25,000"}]) == "₹ 1,75,000"
    assert wf.total_cost_text([{"cost": "50000"}, {"cost": ""}]) == ""
    assert wf.total_cost_text([]) == ""


def test_ai_inputs_carry_excel_context_but_no_contacts(fixtures):
    d = edit(draft_for(fixtures, "01_all_in_scope_full"), ai={"sector": "IT", "duration": "30 days"})
    payload = wf.ai_inputs(d, wf.resolve_plan(d))
    seg = next(c for c in payload["custom_annexures"] if "Segmentation" in c["title"])
    assert seg["count_text"].endswith("Network Segmentation categories") and seg["excel_context"]
    text = json.dumps(payload)
    assert "@example.com" not in text and "+91" not in text                      # no personal data to the AI
    assert payload["commercials"]["total_cost"] == ""                            # D18: optional


def test_context_block_is_capped_and_empty_when_nothing_was_answered(fixtures):
    ctx = proposal_excel.parse_scoping_workbook(sample(fixtures, "18_huge_context")).context
    for key in ("network_segmentation", "saas_security", "web_app_vapt"):
        assert len(proposal_excel.context_block_for_type(ctx, key)) <= 6000
    blank = proposal_excel.parse_scoping_workbook(sample(fixtures, "09_minimal")).context
    assert proposal_excel.context_block_for_type(blank, "network_segmentation") == ""


# ============================================================================ draft store
def test_draft_store_roundtrip_expiry_and_id_validation(tmp_path):
    now = [1000.0]
    store = DraftStore(tmp_path, ttl=60, clock=lambda: now[0])
    draft_id = store.create({"client": {"name": "Acme"}})
    assert store.get(draft_id) == {"client": {"name": "Acme"}}
    assert store.update(draft_id, {"client": {"name": "Beta"}}) and store.get(draft_id)["client"]["name"] == "Beta"
    for bad in ("../etc/passwd", "short", "a" * 100, "x/y" + "a" * 30, "", None):
        assert store.get(bad) is None and store.update(bad, {}) is False
    now[0] += 61
    assert store.get(draft_id) is None and not list(tmp_path.glob("*.json"))     # expired file deleted
    other = store.create({"a": 1})
    now[0] += 61
    assert store.purge() == 1 and store.get(other) is None


# ============================================================================ API
@pytest.fixture()
def client(monkeypatch, tmp_path):
    monkeypatch.setenv("RATE_LIMIT_DISABLED", "1")
    limiter.reset()
    monkeypatch.setattr(gen.proposal_router.draft_store, "dir", tmp_path)   # drafts live in the test's temp folder
    return TestClient(gen.app)


def upload(client, path: Path, name: str | None = None):
    return client.post("/upload-excel/", files={"file": (name or path.name, path.read_bytes(),
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def test_upload_returns_the_two_preview_draft(client, fixtures):
    r = upload(client, sample(fixtures, "31_image_example_counts"))
    assert r.status_code == 200
    body = r.json()
    assert re.match(r"^[A-Za-z0-9_-]{24,64}$", body["draft_id"])
    web = next(e for e in body["draft"]["entries"] if e["id"] == "web_app_vapt")
    assert web["count"] == 6 and body["ttl_seconds"] == 3600
    assert body["draft"]["source_filename"] == "31_image_example_counts.xlsx"


@pytest.mark.parametrize("name, status, fragment", [
    ("23_wrong_sheet_name", 400, "scoping sheet"),
    ("24_missing_sheet", 400, "scoping sheet"),
    ("25_quantity_header_changed", 400, "headers were changed"),
    ("26_empty_workbook", 400, "scoping sheet"),
    ("27_csv_renamed_xlsx", 400, "not a valid .xlsx"),
    ("28_oversize", 413, "larger than 5 MB"),
    ("29_corrupt_zip", 400, "not a valid .xlsx"),
    ("30_all_types_No", 400, "No assessment type"),
])
def test_bad_uploads_are_rejected_with_plain_messages(client, fixtures, name, status, fragment):
    r = upload(client, sample(fixtures, name))
    assert r.status_code == status and fragment in r.json()["detail"]


def test_upload_rejects_other_extensions(client, fixtures):
    r = upload(client, sample(fixtures, "09_minimal"), name="scoping.xlsm")
    assert r.status_code == 400 and "Only .xlsx" in r.json()["detail"]


def test_uploaded_workbook_is_not_kept(client, fixtures, tmp_path):
    import glob, tempfile
    before = set(glob.glob(str(Path(tempfile.gettempdir()) / "css_scoping_*")))
    assert upload(client, sample(fixtures, "09_minimal")).status_code == 200
    assert set(glob.glob(str(Path(tempfile.gettempdir()) / "css_scoping_*"))) == before


def test_draft_get_put_delete_and_unknown_ids(client, fixtures):
    draft_id = upload(client, sample(fixtures, "31_image_example_counts")).json()["draft_id"]
    got = client.get(f"/draft/{draft_id}").json()
    got["draft"]["entries"][0]["frequency"] = "Quarterly"
    put = client.put(f"/draft/{draft_id}", json=got["draft"])
    assert put.status_code == 200 and put.json()["draft"]["entries"][0]["frequency"] == "Quarterly"
    assert client.get(f"/draft/{draft_id}").json()["draft"]["entries"][0]["frequency"] == "Quarterly"
    assert client.delete(f"/draft/{draft_id}").json() == {"deleted": True}
    assert client.get(f"/draft/{draft_id}").status_code == 404
    for bad in ("nope", "../../etc/passwd", "a" * 30):
        assert client.get(f"/draft/{bad}").status_code == 404


def test_generate_is_blocked_until_the_issues_are_resolved(client, fixtures):
    draft_id = upload(client, sample(fixtures, "11_placeholders_untouched")).json()["draft_id"]   # no reference no.
    r = client.post(f"/draft/{draft_id}/generate/")
    assert r.status_code == 400 and "Reference number is required" in r.json()["detail"]


def test_excel_page_is_served_and_has_both_previews(client):
    r = client.get("/excel/")
    assert r.status_code == 200 and "Preview 1" in r.text and "Preview 2" in r.text
    assert r.headers["cache-control"] == "no-store"


def test_manual_form_links_to_the_excel_page(client):
    assert 'href="excel/"' in client.get("/").text


def test_upload_is_rate_limited(monkeypatch, fixtures):
    monkeypatch.delenv("RATE_LIMIT_DISABLED", raising=False)
    monkeypatch.setenv("RATE_LIMIT_UPLOAD", "2/60")
    limiter.reset()
    c = TestClient(gen.app)
    path = sample(fixtures, "09_minimal")
    codes = [upload(c, path).status_code for _ in range(3)]
    assert codes == [200, 200, 429]
    r = upload(c, path)
    assert r.status_code == 429 and int(r.headers["retry-after"]) > 0
    limiter.reset()


def test_cors_does_not_allow_credentials(client):
    r = client.get("/catalog/", headers={"Origin": "https://evil.example"})
    assert "access-control-allow-credentials" not in r.headers


# ============================================================================ generation (no AI)
def _full_draft_id(client, fixtures, name, **edits):
    draft_id = upload(client, sample(fixtures, name)).json()["draft_id"]
    if edits:
        r = client.put(f"/draft/{draft_id}", json=edits)
        assert r.status_code == 200, r.text
    return draft_id


def _doc(response) -> Document:
    assert response.status_code == 200, response.text[:400]
    return Document(io.BytesIO(response.content))


def _field_lines(doc, prefix):
    return [p.text.strip() for p in doc.paragraphs if p.text.strip().startswith(prefix)]


def _pricing_rows(doc):
    tbl = gen._find_assessment_pricing_table(doc.element)
    return [[("".join(t.text or "" for t in tc.iter(qn("w:t")))) for tc in tr.findall(qn("w:tc"))]
            for tr in tbl.findall(qn("w:tr"))]


def test_known_only_generation_end_to_end(client, fixtures):
    draft_id = _full_draft_id(client, fixtures, "02_known_only")
    r = client.post(f"/draft/{draft_id}/generate/")
    doc = _doc(r)
    headings = [p.text.strip() for p in doc.paragraphs if re.match(r"^Annexure \d+:", p.text.strip())]
    assert [int(re.match(r"Annexure (\d+)", h).group(1)) for h in headings] == list(range(1, len(headings) + 1))
    counts = _field_lines(doc, "Count:")
    assert len(counts) == len(headings) == 5
    assert any("Web Application categories" in c for c in counts)
    rows = _pricing_rows(doc)
    assert rows[0][-1] == "Cost Per Annexure" and rows[-1][0] == "Total"
    assert len(rows) == 1 + 5 + 1                                           # header + five annexures + total
    assert proposal_lint.lint_docx(doc).ok
    assert r.headers["x-proposal-lint"].startswith("0 error(s)")


def test_numbers_in_count_lines_and_pricing_match_the_excel(client, fixtures):
    parsed = proposal_excel.parse_scoping_workbook(sample(fixtures, "05_external_and_internal_both"))
    draft_id = _full_draft_id(client, fixtures, "05_external_and_internal_both",
                              entries=[{"id": "external_network_pt", "frequency": "Quarterly"}])
    doc = _doc(client.post(f"/draft/{draft_id}/generate/"))
    assert _field_lines(doc, "Count:") == ["Count: 6 External Infrastructure categories, 12 Internal Infrastructure categories"]
    assert _field_lines(doc, "Frequency:") == ["Frequency: Quarterly"]
    rows = _pricing_rows(doc)
    assert [r[1:4] for r in rows[1:-1]] == [
        ["External Network Penetration Testing", "25", "Quarterly"],
        ["Internal Network Penetration Testing", "90", "Quarterly"]]
    assert {e.quantity for e in parsed.included} == {25, 90}


def test_prices_are_formatted_and_totalled_in_indian_style(client, fixtures):
    draft_id = upload(client, sample(fixtures, "05_external_and_internal_both")).json()["draft_id"]
    pricing = client.get(f"/draft/{draft_id}").json()["draft"]["pricing"]
    pricing[0]["cost"], pricing[1]["cost"] = "50000", "1,25,000"
    doc = _doc(client.post(f"/draft/{draft_id}/generate/", data={"edits_json": json.dumps({"pricing": pricing})}))
    rows = _pricing_rows(doc)
    assert [r[-1] for r in rows[1:-1]] == ["₹ 50,000", "₹ 1,25,000"]
    assert rows[-1][1] == "₹ 1,75,000 /- * + 18% GST"


def test_blank_prices_keep_the_placeholder_total(client, fixtures):
    draft_id = _full_draft_id(client, fixtures, "09_minimal")
    rows = _pricing_rows(_doc(client.post(f"/draft/{draft_id}/generate/")))
    assert rows[-1][1] == "₹ /- * + 18% GST" and rows[1][-1] == ""


def test_removed_pricing_row_is_absent_but_the_annexure_stays(client, fixtures):
    draft_id = upload(client, sample(fixtures, "05_external_and_internal_both")).json()["draft_id"]
    pricing = client.get(f"/draft/{draft_id}").json()["draft"]["pricing"][:1]
    doc = _doc(client.post(f"/draft/{draft_id}/generate/", data={"edits_json": json.dumps({"pricing": pricing})}))
    assert len(_pricing_rows(doc)) == 3                                          # header + 1 + total
    assert len(_field_lines(doc, "Count:")) == 1


def test_zero_quantity_flow_confirm_then_generate(client, fixtures):
    draft_id = upload(client, sample(fixtures, "10_zero_quantity")).json()["draft_id"]
    assert client.post(f"/draft/{draft_id}/generate/").status_code == 400
    view = client.get(f"/draft/{draft_id}").json()
    confirm = [{"id": e["id"], "confirmed": True} for e in view["draft"]["entries"] if e["include"]]
    ok = client.put(f"/draft/{draft_id}", json={"entries": confirm}).json()
    assert ok["issues"] == []
    doc = _doc(client.post(f"/draft/{draft_id}/generate/"))
    assert any(c.startswith("Count: 0 ") for c in _field_lines(doc, "Count:"))   # confirmed as is


def test_manual_toggle_adds_an_annexure_without_excel_rows(client, fixtures):
    draft_id = _full_draft_id(client, fixtures, "09_minimal", entries=[{"id": "post_breach", "include": True}])
    doc = _doc(client.post(f"/draft/{draft_id}/generate/"))
    titles = [p.text.strip() for p in doc.paragraphs if re.match(r"^Annexure \d+:", p.text.strip())]
    assert titles == ["Annexure 1: Post Breach Scenario Assessment", "Annexure 2: Web Application Penetration Testing Assessment"]
    assert "Count: 1 Assessment categories" in _field_lines(doc, "Count:")


def test_user_chosen_order_is_respected_and_renumbered(client, fixtures):
    draft_id = upload(client, sample(fixtures, "02_known_only")).json()["draft_id"]
    ids = [e["id"] for e in client.get(f"/draft/{draft_id}").json()["draft"]["entries"]]   # the page sends them all
    reordered = [{"id": i} for i in reversed(ids)]
    doc = _doc(client.post(f"/draft/{draft_id}/generate/", data={"edits_json": json.dumps({"entries": reordered})}))
    headings = [p.text.strip() for p in doc.paragraphs if re.match(r"^Annexure \d+:", p.text.strip())]
    assert [int(re.match(r"Annexure (\d+)", h).group(1)) for h in headings] == [1, 2, 3, 4, 5]
    assert "Source Code Review" in headings[0]
    assert proposal_lint.lint_docx(doc).ok


def test_manual_flow_still_works(client):
    r = client.post("/generate/", data={"company_name": "Manual Co", "proposal_ref": "EN-M",
                    "annexure_order_json": json.dumps(["web_app_pt", "cloud_pt"])})
    doc = _doc(r)
    assert len([p for p in doc.paragraphs if re.match(r"^Annexure \d+:", p.text.strip())]) == 2
    assert _pricing_rows(doc)[0][-1] == "Cost Per Annexure"
    assert proposal_lint.lint_docx(doc).ok


# ============================================================================ generation with AI (mocked)
AI_TEXT = """Annexure 77: Network Segmentation / VLAN Testing

Sub-Activity:
• Assess VLAN isolation and inter-zone access control across the agreed segments.

Methodologies:
• Manual testing against the documented segmentation policy
o Review of ACL rule sets and routing between zones

Frequency: One time with Revalidation

Count: 999 Wrong Things

Deliverables:
• Technical report with evidence for every tested zone boundary and recommended fixes.

Expectation:
• Client provides network diagrams and read-only access to the firewall rule sets.

Objectives:
• Confirm that the production zone cannot be reached from lower trust zones.

Out of Scope:
• Denial of service testing and physical access to network equipment.

Requirements for Starting the Assessment:
• Signed authorization, named contacts and a confirmed testing window.

Comments:
• Findings are reported with severity and a remediation owner for each item.
"""
BAD_AI_TEXT = "Annexure 1: X\n\nSub-Activity:\n• too short and missing the other required sections."


class FakeOpenAI:
    """Replaces openai.AsyncOpenAI: returns the queued answers, records every request."""
    answers: list[str] = []
    calls: list[dict] = []

    def __init__(self, *a, **k):
        async def create(**kwargs):
            FakeOpenAI.calls.append(kwargs)
            text = FakeOpenAI.answers.pop(0) if len(FakeOpenAI.answers) > 1 else FakeOpenAI.answers[0]
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=text))])
        self.chat = SimpleNamespace(completions=SimpleNamespace(create=create))


@pytest.fixture()
def fake_ai(monkeypatch):
    import openai
    FakeOpenAI.answers, FakeOpenAI.calls = [AI_TEXT], []
    monkeypatch.setattr(openai, "AsyncOpenAI", FakeOpenAI)
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test-not-real")
    return FakeOpenAI


AI_FIELDS = {"ai": {"sector": "Banking", "duration": "30 days", "engagement_summary": "Segmentation review for a bank."}}


def test_mixed_flow_with_an_ai_annexure(client, fixtures, fake_ai):
    draft_id = _full_draft_id(client, fixtures, "04_mixed_known_and_custom", **AI_FIELDS)
    r = client.post(f"/draft/{draft_id}/generate/")
    doc = _doc(r)
    headings = [p.text.strip() for p in doc.paragraphs if re.match(r"^Annexure \d+:", p.text.strip())]
    nums = [int(re.match(r"Annexure (\d+)", h).group(1)) for h in headings]
    assert nums == list(range(1, len(headings) + 1))                       # AI heading number (77) corrected
    seg = next(h for h in headings if "Network Segmentation" in h)
    assert not seg.startswith("Annexure 77")
    counts = _field_lines(doc, "Count:")
    assert "Count: 999 Wrong Things" not in counts                         # AI's Count line replaced
    assert any(c.endswith("Network Segmentation categories") for c in counts)
    assert len(fake_ai.calls) == 2                                          # Network Segmentation + Red Team
    assert proposal_lint.lint_docx(doc).ok
    user_msg = fake_ai.calls[0]["messages"][1]["content"]
    assert "Count for this annexure:" in user_msg and "Not yet specified" in user_msg    # D18 + count


def test_custom_only_flow(client, fixtures, fake_ai):
    draft_id = _full_draft_id(client, fixtures, "03_custom_only", **AI_FIELDS)
    doc = _doc(client.post(f"/draft/{draft_id}/generate/"))
    headings = [p.text.strip() for p in doc.paragraphs if re.match(r"^Annexure \d+:", p.text.strip())]
    assert len(headings) == 5 and len(fake_ai.calls) == 5
    assert len(_pricing_rows(doc)) == 5 + 2
    assert proposal_lint.lint_docx(doc).ok


def test_ai_total_cost_is_derived_from_prices(client, fixtures, fake_ai):
    draft_id = upload(client, sample(fixtures, "04_mixed_known_and_custom")).json()["draft_id"]
    pricing = client.get(f"/draft/{draft_id}").json()["draft"]["pricing"]
    for r in pricing:
        r["cost"] = "10000"
    body = {"pricing": pricing, **AI_FIELDS}
    _doc(client.post(f"/draft/{draft_id}/generate/", data={"edits_json": json.dumps(body)}))
    assert "₹ 50,000" in fake_ai.calls[0]["messages"][1]["content"]


def test_ai_retry_succeeds_on_the_second_answer(client, fixtures, fake_ai):
    fake_ai.answers = [BAD_AI_TEXT, AI_TEXT]
    draft_id = _full_draft_id(client, fixtures, "04_mixed_known_and_custom", entries=[{"id": "red_team", "include": False}], **AI_FIELDS)
    r = client.post(f"/draft/{draft_id}/generate/")
    assert r.status_code == 200
    assert len(fake_ai.calls) == 2
    retry_msgs = fake_ai.calls[1]["messages"]
    assert retry_msgs[-1]["role"] == "user" and "rejected" in retry_msgs[-1]["content"]
    assert retry_msgs[-2]["content"] == BAD_AI_TEXT


def test_ai_failure_after_one_retry_is_a_clear_error(client, fixtures, fake_ai):
    fake_ai.answers = [BAD_AI_TEXT]                                         # every answer is bad
    draft_id = _full_draft_id(client, fixtures, "04_mixed_known_and_custom", entries=[{"id": "red_team", "include": False}], **AI_FIELDS)
    r = client.post(f"/draft/{draft_id}/generate/")
    assert r.status_code == 502 and "retry" in r.json()["detail"].lower()
    assert len(fake_ai.calls) == 2                                          # exactly one retry, no loop


def test_missing_openai_key_is_reported(client, fixtures, monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "")
    monkeypatch.setattr(gen, "HERE", Path(fixtures))                         # no .env here, so no key can be loaded
    draft_id = _full_draft_id(client, fixtures, "04_mixed_known_and_custom", **AI_FIELDS)
    r = client.post(f"/draft/{draft_id}/generate/")
    assert r.status_code == 500 and "not configured" in r.json()["detail"]
