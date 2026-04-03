"""
Type 5: Nessus CSV to Excel Report Generator
Accepts a Nessus vulnerability scanner CSV export and produces a formatted
Excel workbook with two sheets -- Summary and Vuln.

Hardened for production: handles encoding quirks, illegal characters,
missing columns, oversized cells, malformed rows, and large files.
"""

from fastapi import FastAPI, UploadFile, File, HTTPException, Request
from fastapi.responses import FileResponse, HTMLResponse
from fastapi.middleware.cors import CORSMiddleware
import csv
import io
import os
import re
import tempfile
import traceback
import logging
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("type5")

app = FastAPI()
# Session cookie is owned only by Report-Generator-IP-main/app.py (no nested SessionMiddleware here).

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

SEVERITY_ORDER = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3, "Informational": 4}
VALID_RISKS = {"Critical", "High", "Medium", "Low", "None", ""}

REQUIRED_COLUMNS = {"Name", "Risk", "IP Address"}
OPTIONAL_COLUMNS = {"CVE", "CVSS", "CVSS3 Base Score", "Synopsis", "Description",
                    "Solution", "See Also", "Host", "Protocol", "Port", "Plugin ID"}

MAX_CELL_LENGTH = 32_000  # Excel limit is 32,767; leave margin
MAX_FILE_SIZE = 100 * 1024 * 1024  # 100 MB

SUMMARY_COLORS = {
    "location": "7030A0",
    "Critical": "C00000",
    "High": "FF0000",
    "Medium": "FFC000",
    "Low": "00B050",
    "Informational": "0070C0",
    "Total": "7030A0",
}

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

_ILLEGAL_XML_RE = re.compile(
    r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f\ud800-\udfff\ufdd0-\ufdef\ufffe\uffff]"
)


# ---------------------------------------------------------------------------
# Utility helpers
# ---------------------------------------------------------------------------

def clean_for_excel(text: str) -> str:
    """Remove characters illegal in XML/Excel and enforce cell length limit."""
    if not text:
        return ""
    cleaned = _ILLEGAL_XML_RE.sub("", text)
    if len(cleaned) > MAX_CELL_LENGTH:
        cleaned = cleaned[:MAX_CELL_LENGTH] + "... [truncated]"
    return cleaned


def safe_str(value) -> str:
    """Safely convert any value to a stripped string."""
    if value is None:
        return ""
    return str(value).strip()


def safe_float(value, default: float = 0.0) -> float:
    """Safely convert to float."""
    try:
        return float(value)
    except (ValueError, TypeError):
        return default


def decode_csv_bytes(raw: bytes) -> str:
    """Try multiple encodings to decode CSV bytes.
    Order: utf-8-sig (handles BOM), utf-8, cp1252 (Windows), latin-1 (never fails).
    """
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return raw.decode("latin-1", errors="replace")


def detect_delimiter(first_line: str) -> str:
    """Auto-detect CSV delimiter from the header line."""
    for delim in (",", ";", "\t", "|"):
        if delim in first_line:
            count = first_line.count(delim)
            if count >= 3:
                return delim
    return ","


def extract_location(filename: str) -> str:
    """Extract location name from CSV filename (e.g. Chennai_VLAN.csv -> Chennai)."""
    if not filename:
        return "Unknown"
    stem = os.path.splitext(os.path.basename(filename))[0]
    parts = re.split(r"[_\-\s]+", stem)
    return parts[0] if parts else stem


def map_risk(risk_value: str) -> str:
    """Map CSV Risk value to a standard severity label."""
    val = safe_str(risk_value)
    if not val or val.lower() == "none":
        return "Informational"
    # Case-insensitive match
    for standard in ("Critical", "High", "Medium", "Low"):
        if val.lower() == standard.lower():
            return standard
    return "Informational"


def validate_columns(fieldnames: list[str]) -> tuple[bool, list[str]]:
    """Check that the CSV has the minimum required columns.
    Returns (is_valid, list_of_missing_columns).
    """
    if not fieldnames:
        return False, list(REQUIRED_COLUMNS)
    present = {f.strip() for f in fieldnames if f}
    missing = [c for c in REQUIRED_COLUMNS if c not in present]

    # Also accept 'Host' as fallback for 'IP Address'
    if "IP Address" in missing and "Host" in present:
        missing.remove("IP Address")

    return len(missing) == 0, missing


# ---------------------------------------------------------------------------
# Core processing
# ---------------------------------------------------------------------------

def parse_and_group(csv_text: str):
    """Parse CSV text and group rows by vulnerability Name.

    Returns a list of dicts, each representing a unique vulnerability with
    aggregated IPs and metadata, sorted by severity then name.
    """
    lines = csv_text.splitlines(keepends=True)
    if not lines:
        return []

    delimiter = detect_delimiter(lines[0])
    reader = csv.DictReader(io.StringIO(csv_text), delimiter=delimiter)

    if not reader.fieldnames:
        return []

    is_valid, missing = validate_columns(list(reader.fieldnames))
    if not is_valid:
        raise ValueError(f"CSV is missing required columns: {', '.join(missing)}")

    has_ip_address = "IP Address" in reader.fieldnames
    vuln_map: dict[str, list[dict]] = defaultdict(list)
    skipped = 0

    for row in reader:
        name = safe_str(row.get("Name"))
        if not name:
            skipped += 1
            continue
        vuln_map[name].append(row)

    if skipped > 0:
        logger.info("Skipped %d rows with empty Name", skipped)

    grouped: list[dict] = []
    for name, rows in vuln_map.items():
        risk_raw = safe_str(rows[0].get("Risk"))
        severity = map_risk(risk_raw)

        ip_field = "IP Address" if has_ip_address else "Host"
        unique_ips = sorted(
            {safe_str(r.get(ip_field)) for r in rows} - {""}
        )
        cves = sorted(
            {safe_str(r.get("CVE")) for r in rows} - {""}
        )

        cvss_values = []
        for r in rows:
            for field in ("CVSS3 Base Score", "CVSS"):
                v = safe_float(r.get(field))
                if v > 0:
                    cvss_values.append(v)
        max_cvss = max(cvss_values) if cvss_values else 0.0

        first = rows[0]
        solution = safe_str(first.get("Solution"))
        if solution.upper() in ("N/A", "NA", "NONE", ""):
            solution = ""

        grouped.append(
            {
                "name": clean_for_excel(name),
                "severity": severity,
                "cves": clean_for_excel(", ".join(cves)),
                "cvss": max_cvss,
                "ips": "\n".join(unique_ips),
                "ip_count": len(unique_ips),
                "synopsis": clean_for_excel(safe_str(first.get("Synopsis"))),
                "description": clean_for_excel(safe_str(first.get("Description"))),
                "solution": clean_for_excel(solution),
                "see_also": clean_for_excel(safe_str(first.get("See Also"))),
            }
        )

    grouped.sort(key=lambda v: (SEVERITY_ORDER.get(v["severity"], 99), v["name"]))
    return grouped


# ---------------------------------------------------------------------------
# Excel generation
# ---------------------------------------------------------------------------

def _safe_write_cell(ws, row, col, value, **style_kwargs):
    """Write a value to a cell, catching any openpyxl errors."""
    try:
        if isinstance(value, str):
            value = clean_for_excel(value)
        cell = ws.cell(row=row, column=col, value=value)
    except Exception:
        cell = ws.cell(row=row, column=col, value=str(value)[:MAX_CELL_LENGTH])

    for attr, val in style_kwargs.items():
        try:
            setattr(cell, attr, val)
        except Exception:
            pass
    return cell


def build_summary_sheet(ws, location: str, vulns: list[dict]):
    """Build the Summary sheet matching the template format."""
    counts = {s: 0 for s in SEVERITY_ORDER}
    for v in vulns:
        sev = v["severity"]
        counts[sev] = counts.get(sev, 0) + 1
    total = sum(counts.values())

    headers = ["Critical", "High", "Medium", "Low", "Informational", "Total"]
    values = [counts.get(h, 0) for h in headers[:-1]] + [total]

    ws.merge_cells("A1:A3")
    loc_cell = ws["A1"]
    loc_cell.value = clean_for_excel(location)
    loc_cell.font = Font(bold=True, size=11, color="FFFFFF")
    loc_cell.fill = PatternFill(start_color=SUMMARY_COLORS["location"], fill_type="solid")
    loc_cell.alignment = Alignment(horizontal="center", vertical="center")
    loc_cell.border = THIN_BORDER

    for col_idx, header in enumerate(headers, start=2):
        col_letter = get_column_letter(col_idx)
        ws.merge_cells(f"{col_letter}1:{col_letter}2")
        cell = ws.cell(row=1, column=col_idx)
        cell.value = header
        cell.font = Font(bold=True, size=11, color="FFFFFF")
        color_key = header if header in SUMMARY_COLORS else "Total"
        cell.fill = PatternFill(start_color=SUMMARY_COLORS[color_key], fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

        val_cell = ws.cell(row=3, column=col_idx)
        val_cell.value = values[col_idx - 2]
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
        val_cell.border = THIN_BORDER

    for row_cells in ws.iter_rows(min_row=2, max_row=3, min_col=1, max_col=1):
        for c in row_cells:
            c.border = THIN_BORDER

    ws.column_dimensions["A"].width = 18
    for col_idx in range(2, 8):
        ws.column_dimensions[get_column_letter(col_idx)].width = 16


def build_vuln_sheet(ws, location: str, vulns: list[dict]):
    """Build the Vuln sheet matching the template format."""
    col_headers = [
        "Sr. No.",
        "Name of Vulnerability",
        "Risk(Severity)",
        "CVE",
        "CVSS",
        "Affected Ips",
        "Audit Observation",
        "Impact",
        "Recommendation",
        "Reference Link",
        "Revalidation",
    ]
    num_cols = len(col_headers)

    # Row 1: Sr. No. label in A1, rest merged blank
    _safe_write_cell(ws, 1, 1, "Sr. No.",
                     font=Font(bold=True, size=12), border=THIN_BORDER)
    if num_cols > 1:
        ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=num_cols - 1)
        for c in range(2, num_cols):
            ws.cell(row=1, column=c).border = THIN_BORDER

    # Row 2: Location label in A2, location value merged across
    _safe_write_cell(ws, 2, 1, "Location",
                     font=Font(bold=True, size=12), border=THIN_BORDER)
    if num_cols > 1:
        ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=num_cols - 1)
        _safe_write_cell(ws, 2, 2, f"{location} Private IP",
                         font=Font(size=12), border=THIN_BORDER)

    # Row 3: Column headers
    header_fill = PatternFill(start_color="4472C4", fill_type="solid")
    for ci, header in enumerate(col_headers, start=1):
        _safe_write_cell(ws, 3, ci, header,
                         font=Font(bold=True, size=12, color="FFFFFF"),
                         fill=header_fill,
                         alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
                         border=THIN_BORDER)

    # Data rows starting at row 4
    severity_fills = {
        "Critical": PatternFill(start_color="C00000", fill_type="solid"),
        "High": PatternFill(start_color="FF0000", fill_type="solid"),
        "Medium": PatternFill(start_color="FFC000", fill_type="solid"),
        "Low": PatternFill(start_color="00B050", fill_type="solid"),
        "Informational": PatternFill(start_color="0070C0", fill_type="solid"),
    }
    severity_fonts = {
        "Critical": Font(size=11, color="FFFFFF", bold=True),
        "High": Font(size=11, color="FFFFFF", bold=True),
        "Medium": Font(size=11, color="000000", bold=True),
        "Low": Font(size=11, color="FFFFFF", bold=True),
        "Informational": Font(size=11, color="FFFFFF", bold=True),
    }

    default_font = Font(size=11)
    wrap_top = Alignment(vertical="top", wrap_text=True)
    col_widths = [8, 35, 16, 20, 8, 22, 40, 40, 40, 30, 14]
    LINE_HEIGHT_PT = 16
    MAX_ROW_HEIGHT = 409  # Excel maximum is 409.5pt

    for idx, v in enumerate(vulns, start=1):
        row = idx + 3
        row_values = [
            idx,
            v["name"],
            v["severity"],
            v["cves"],
            v["cvss"] if v["cvss"] > 0 else "",
            v["ips"],
            v["synopsis"],
            v["description"],
            v["solution"],
            v["see_also"],
            "",
        ]

        max_lines = 1
        for ci, val in enumerate(row_values, start=1):
            _safe_write_cell(ws, row, ci, val,
                             font=default_font,
                             alignment=wrap_top,
                             border=THIN_BORDER)
            if isinstance(val, str) and val:
                col_w = col_widths[ci - 1] if ci <= len(col_widths) else 14
                # SharePoint renders wider than desktop Excel, so use
                # a conservative estimate: ~0.85 chars per width unit
                chars_per_line = max(int(col_w * 0.85), 1)
                lines = 0
                for segment in val.split("\n"):
                    seg_len = len(segment)
                    lines += max(1, -(-seg_len // chars_per_line))
                max_lines = max(max_lines, lines)

        # Add 1 extra line of padding so SharePoint never clips
        computed = LINE_HEIGHT_PT * (max_lines + 1)
        ws.row_dimensions[row].height = min(max(computed, LINE_HEIGHT_PT * 2), MAX_ROW_HEIGHT)

        sev = v["severity"]
        if sev in severity_fills:
            sev_cell = ws.cell(row=row, column=3)
            sev_cell.fill = severity_fills[sev]
            sev_cell.font = severity_fonts[sev]
            sev_cell.alignment = Alignment(horizontal="center", vertical="center")
    for ci, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = width


def generate_excel(location: str, vulns: list[dict]) -> str:
    """Create the Excel workbook and return the path to the temp file."""
    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Summary"
    build_summary_sheet(ws_summary, location, vulns)

    ws_vuln = wb.create_sheet("Vuln")
    build_vuln_sheet(ws_vuln, location, vulns)

    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        wb.save(path)
    except Exception:
        if os.path.exists(path):
            os.unlink(path)
        raise
    return path


# ---------------------------------------------------------------------------
# FastAPI endpoints
# ---------------------------------------------------------------------------

@app.get("/", response_class=HTMLResponse)
async def root():
    html_content = """
<!DOCTYPE html>
<html>
<head>
    <title>Nessus CSV to Excel - VAPT Report Generator</title>
    <script src="/static/js/csrf.js"></script>
    <style>
        * { box-sizing: border-box; margin: 0; padding: 0; }
        body {
            font-family: Arial, sans-serif;
            background: #1a1a2e;
            color: #e0e0e0;
            min-height: 100vh;
            display: flex;
            justify-content: center;
            align-items: flex-start;
            padding: 40px 20px;
        }
        .container {
            background: #16213e;
            border-radius: 12px;
            padding: 40px;
            max-width: 700px;
            width: 100%;
            box-shadow: 0 8px 32px rgba(0,0,0,0.3);
        }
        h1 {
            color: #8a2be2;
            margin-bottom: 8px;
            font-size: 1.6em;
        }
        .subtitle {
            color: #999;
            margin-bottom: 30px;
            font-size: 0.95em;
        }
        .form-group {
            margin-bottom: 22px;
        }
        label {
            display: block;
            font-weight: bold;
            margin-bottom: 8px;
            color: #ccc;
        }
        input[type="file"] {
            width: 100%;
            padding: 12px;
            border: 2px dashed #6923d0;
            border-radius: 8px;
            background: rgba(105,35,208,0.08);
            color: #e0e0e0;
            cursor: pointer;
            transition: border-color 0.2s;
        }
        input[type="file"]:hover { border-color: #8a2be2; }
        .btn {
            background: #6923d0;
            color: #fff;
            border: none;
            padding: 14px 32px;
            border-radius: 8px;
            font-size: 1.05em;
            font-weight: bold;
            cursor: pointer;
            width: 100%;
            transition: background 0.2s;
        }
        .btn:hover { background: #5a1fb8; }
        .btn:disabled { background: #444; cursor: not-allowed; }
        .info-box {
            background: rgba(0,0,0,0.2);
            border-left: 4px solid #6923d0;
            padding: 14px 18px;
            border-radius: 4px;
            margin-bottom: 24px;
            font-size: 0.92em;
            line-height: 1.6;
        }
        .info-box strong { color: #8a2be2; }
        #status {
            margin-top: 18px;
            padding: 14px;
            border-radius: 8px;
            display: none;
            font-weight: bold;
            text-align: center;
        }
        .status-processing { background: rgba(105,35,208,0.15); color: #b388ff; display: block !important; }
        .status-success { background: rgba(0,176,80,0.15); color: #69f0ae; display: block !important; }
        .status-error { background: rgba(255,68,68,0.15); color: #ff5252; display: block !important; }
        a.back { color: #8a2be2; text-decoration: none; display: inline-block; margin-bottom: 20px; }
        a.back:hover { text-decoration: underline; }
    </style>
</head>
<body>
    <div class="container">
        <a class="back" href="/report_formats.html">&larr; Back to Report Formats</a>
        <h1>Nessus CSV to Excel</h1>
        <p class="subtitle">Upload a Nessus vulnerability scanner CSV export to generate a formatted Excel report.</p>

        <div class="info-box">
            <strong>How it works:</strong><br>
            1. Upload the <code>.csv</code> file exported from Nessus<br>
            2. Location is auto-detected from the filename<br>
            3. Vulnerabilities are grouped by name with affected IPs aggregated<br>
            4. Output Excel has two sheets: <strong>Summary</strong> (severity counts) and <strong>Vuln</strong> (detailed findings)
        </div>

        <form id="uploadForm" enctype="multipart/form-data">
            <div class="form-group">
                <label for="csvFile">Upload Nessus CSV File</label>
                <input type="file" id="csvFile" name="file" accept=".csv" required>
            </div>
            <button type="submit" class="btn" id="submitBtn">Generate Excel Report</button>
        </form>

        <div id="status"></div>
    </div>

    <script>
    document.getElementById('uploadForm').addEventListener('submit', async function(e) {
        e.preventDefault();
        const fileInput = document.getElementById('csvFile');
        const status = document.getElementById('status');
        const btn = document.getElementById('submitBtn');

        if (!fileInput.files[0]) { alert('Please select a CSV file.'); return; }

        const fileSizeMB = fileInput.files[0].size / 1024 / 1024;
        if (fileSizeMB > 100) {
            alert('File is too large (max 100 MB). Got ' + fileSizeMB.toFixed(1) + ' MB.');
            return;
        }

        btn.disabled = true;
        status.className = 'status-processing';
        status.textContent = 'Processing CSV (' + fileSizeMB.toFixed(1) + ' MB) and generating Excel...';

        const formData = new FormData();
        formData.append('file', fileInput.files[0]);

        try {
            const resp = await fetch('/type5/generate-report/', {
                method: 'POST',
                body: formData,
                credentials: 'same-origin'
            });
            if (resp.status === 401) {
                throw new Error('Session expired. Please <a href="/auth/test-login">log in again</a>.');
            }
            if (!resp.ok) {
                let errMsg = 'Generation failed';
                try {
                    const err = await resp.json();
                    errMsg = err.detail || errMsg;
                } catch (_) {
                    errMsg = 'Server error (HTTP ' + resp.status + '). The file may be too large or malformed.';
                }
                throw new Error(errMsg);
            }
            const blob = await resp.blob();
            const cd = resp.headers.get('content-disposition') || '';
            const match = cd.match(/filename="?([^"]+)"?/);
            const filename = match ? match[1] : 'report.xlsx';

            const url = URL.createObjectURL(blob);
            const a = document.createElement('a');
            a.href = url; a.download = filename;
            document.body.appendChild(a); a.click(); a.remove();
            URL.revokeObjectURL(url);

            status.className = 'status-success';
            status.textContent = 'Report generated and downloaded successfully!';
        } catch (err) {
            status.className = 'status-error';
            status.innerHTML = 'Error: ' + err.message;
        } finally {
            btn.disabled = false;
        }
    });
    </script>
</body>
</html>
"""
    return HTMLResponse(content=html_content)


@app.post("/generate-report/")
async def generate_report(file: UploadFile = File(...)):
    # 1. Validate file extension
    if not file.filename or not file.filename.lower().endswith(".csv"):
        raise HTTPException(status_code=400, detail="Please upload a .csv file")

    # 2. Read and check file size
    try:
        raw = await file.read()
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read uploaded file: {e}")

    if len(raw) == 0:
        raise HTTPException(status_code=400, detail="The uploaded CSV file is empty.")

    if len(raw) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=400,
            detail=f"File too large ({len(raw)/1024/1024:.1f} MB). Maximum is {MAX_FILE_SIZE//1024//1024} MB."
        )

    # 3. Decode with encoding fallback chain
    csv_text = decode_csv_bytes(raw)

    # 4. Basic CSV sanity check -- must have at least a header + 1 data row
    line_count = csv_text.count("\n")
    if line_count < 2:
        raise HTTPException(
            status_code=400,
            detail="CSV file appears to have no data rows (only header or empty)."
        )

    # 5. Parse and group
    try:
        vulns = parse_and_group(csv_text)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        logger.error("CSV parsing failed: %s\n%s", e, traceback.format_exc())
        raise HTTPException(
            status_code=500,
            detail=f"Failed to parse CSV: {type(e).__name__}: {e}"
        )

    if not vulns:
        raise HTTPException(
            status_code=400,
            detail="No vulnerabilities found in the CSV. Ensure it has 'Name', 'Risk', and 'IP Address' columns."
        )

    # 6. Generate Excel
    location = extract_location(file.filename)
    try:
        output_path = generate_excel(location, vulns)
    except Exception as e:
        logger.error("Excel generation failed: %s\n%s", e, traceback.format_exc())
        raise HTTPException(
            status_code=500,
            detail=f"Failed to generate Excel: {type(e).__name__}: {e}"
        )

    output_name = f"{location}_VAPT_Report.xlsx"

    return FileResponse(
        path=output_path,
        filename=output_name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        background=None,
    )


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8005)
