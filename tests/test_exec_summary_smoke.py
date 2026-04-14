"""Smoke tests for Executive Summary generator (no live OpenAI calls)."""

from __future__ import annotations

import io
import os
import sys
import tempfile
import unittest
import uuid
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


def _sample_ai_payload() -> dict:
    long_para = (
        "This assessment reviewed the agreed scope and identified themes that warrant executive attention. "
        "The engagement followed industry-recognized practices and focused on practical risk to the business. "
        "Findings were triaged by severity to support remediation planning and governance discussions. "
        "The organization demonstrated willingness to collaborate during validation windows and evidence review. "
        "Residual risk should be tracked until remediation and revalidation are complete across environments. "
        "Leadership should ensure ownership for fixes, timelines, and communication with stakeholders. "
        "The following sections summarise observations and recommendations in plain language for decision-makers."
    )
    intro = (
        "Leadership should treat security remediation as a sustained program rather than a one-time activity. "
        "Governance and periodic validation remain important as systems and threats evolve over time."
    )
    p2 = (
        "Post-remediation validation will confirm whether fixes hold under realistic conditions. "
        "Until then, some exposure may remain where compensating controls are not fully effective."
    )
    conc = (
        "In summary, the assessment provides a structured view of risk and practical next steps for the organisation. "
        "We appreciate the security team support during the engagement and stand ready to assist with follow-up.\n\n"
        "Closing remarks: track remediation and validate fixes in production-like conditions."
    )
    item = {
        "title": "Sample theme",
        "description": "This theme reflects findings observed during testing and aligns with the data provided.",
        "requirement": "High",
        "business_impact": "Operational Capabilities",
    }
    return {
        "executive_narrative": long_para + "\n\n" + long_para,
        "overall_recommendation_intro": intro,
        "phase1_observations": [
            {"title": "A", "description": "Desc one. More detail.", "threat_perception": "High"},
            {"title": "B", "description": "Desc two. More detail.", "threat_perception": "Moderate"},
            {"title": "C", "description": "Desc three. More detail.", "threat_perception": "Low"},
        ],
        "phase2_summary": p2,
        "best_practices": [item, item.copy(), {**item, "title": "BP3"}],
        "must_have": [item.copy(), {**item, "title": "M2"}, {**item, "title": "M3"}],
        "should_have": [{**item, "title": "S1"}, {**item, "title": "S2"}],
        "conclusion": conc,
    }


class TestExecSummarySmoke(unittest.TestCase):
    def test_parse_json_normalizes_short_arrays(self):
        """AI sometimes returns too few phase1 rows; parser pads instead of failing."""
        from Automation.backend.exec_summary_prompts import parse_exec_summary_json

        minimal = """
        {
          "executive_narrative": """ + '"' + ("x" * 220) + """",
          "overall_recommendation_intro": "Intro paragraph one. Intro paragraph two with enough chars here.",
          "phase1_observations": [],
          "phase2_summary": "Phase two summary text that is long enough for the minimum length check here.",
          "best_practices": [],
          "must_have": [],
          "should_have": [],
          "conclusion": "Conclusion text that is long enough for the minimum length check to pass validation here."
        }
        """
        out = parse_exec_summary_json(minimal)
        self.assertGreaterEqual(len(out["phase1_observations"]), 3)
        self.assertGreaterEqual(len(out["best_practices"]), 3)
        self.assertGreaterEqual(len(out["must_have"]), 3)
        self.assertGreaterEqual(len(out["should_have"]), 2)

    def test_exec_summary_app_importable(self):
        from Automation.backend import exec_summary  # noqa: F401

        self.assertIsNotNone(exec_summary.app)

    def test_routes_registered(self):
        from Automation.backend.exec_summary import app

        paths = {getattr(r, "path", None) for r in app.routes}
        self.assertIn("/", paths)
        self.assertIn("/generate/", paths)
        self.assertIn("/health", paths)

    def test_build_document_contains_sections(self):
        from Automation.backend.exec_summary import _build_document

        meta = {
            "client_name": "Test Client",
            "org_name": "CyberSmithSECURE",
            "report_date": "2025-04-07",
            "assessment_type": "Web Application VAPT",
            "document_id": "DOC-TEST-001",
        }
        assets = [
            {
                "hostname": "app.example.com",
                "purpose": "App",
                "vapt_status": "Completed",
                "critical": 1,
                "high": 0,
                "medium": 1,
                "low": 0,
                "informational": 0,
                "total": 2,
            }
        ]
        total_counts = {
            "critical": 1,
            "high": 0,
            "medium": 1,
            "low": 0,
            "informational": 0,
            "total": 2,
        }
        doc = _build_document(meta, assets, total_counts, _sample_ai_payload())
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        with zipfile.ZipFile(buf, "r") as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
            # Branded template uses header2.xml for content section (document ID)
            hdr_parts = []
            for name in ("word/header1.xml", "word/header2.xml"):
                if name in zf.namelist():
                    hdr_parts.append(zf.read(name).decode("utf-8"))
            hdr_combined = "\n".join(hdr_parts)
        self.assertIn("Executive Summary of The Assessment", xml)
        self.assertIn("Overall Observation", xml)
        self.assertIn("SUMMARY OF FINDINGS", xml)
        self.assertIn("CONCLUSION:", xml)
        self.assertIn("DOC-TEST-001", hdr_combined)

    def test_build_document_toc_outline_metadata(self):
        """Manual TOC lists section titles in the body; PAGEREF + bookmarks supply page numbers after field update.

        In Word: Enable Editing if needed, then Ctrl+A and F9 so PAGEREF page numbers refresh.
        """
        from Automation.backend.exec_summary import _build_document

        meta = {
            "client_name": "Test Client",
            "org_name": "CyberSmithSECURE",
            "report_date": "2025-04-07",
            "assessment_type": "Web Application VAPT",
            "document_id": "DOC-TOC-001",
        }
        assets = [
            {
                "hostname": "app.example.com",
                "purpose": "App",
                "vapt_status": "Completed",
                "critical": 1,
                "high": 0,
                "medium": 1,
                "low": 0,
                "informational": 0,
                "total": 2,
            }
        ]
        total_counts = {
            "critical": 1,
            "high": 0,
            "medium": 1,
            "low": 0,
            "informational": 0,
            "total": 2,
        }
        doc = _build_document(meta, assets, total_counts, _sample_ai_payload())
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        with zipfile.ZipFile(buf, "r") as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
        # Manual TOC: visible titles + PAGEREF fields (not a single TOC \\o field that stays empty).
        self.assertIn("PAGEREF", xml)
        self.assertIn('w:name="bm_es"', xml)
        self.assertIn("Test Client Risk Assessment Analysis", xml)
        self.assertIn("Best Practises", xml)
        # Built-in Heading 1/2 + outline levels on headings.
        self.assertGreaterEqual(xml.count("Heading1"), 6)
        self.assertGreaterEqual(xml.count("Heading2"), 5)
        outline_count = xml.count("w:outlineLvl")
        self.assertGreaterEqual(outline_count, 11, "Expected outline levels on TOC-eligible headings")
        self.assertGreaterEqual(xml.count("<w:fldChar"), 10)

    def test_parse_minimal_excel(self):
        from openpyxl import Workbook

        from Automation.backend.exec_summary import REQUIRED_HEADERS, _parse_excel_to_payload

        wb = Workbook()
        ws = wb.active
        for i, h in enumerate(REQUIRED_HEADERS, start=1):
            ws.cell(row=1, column=i, value=h)
        ws.cell(row=2, column=1, value="host1")
        ws.cell(row=2, column=12, value="1")
        # Column 21 = "Observation/Vulnerability Title" (see REQUIRED_HEADERS order)
        ws.cell(row=2, column=21, value="SQL injection in login")
        ws.cell(row=2, column=14, value="High")
        ws.cell(row=2, column=4, value=1)
        ws.cell(row=2, column=5, value=0)
        ws.cell(row=2, column=6, value=0)
        ws.cell(row=2, column=7, value=0)
        ws.cell(row=2, column=8, value=0)
        ws.cell(row=2, column=9, value=1)

        path = os.path.join(tempfile.gettempdir(), f"exec_summary_test_{uuid.uuid4().hex}.xlsx")
        try:
            wb.save(path)
            wb.close()
            meta = {"client_name": "X", "assessment_type": "Test"}
            assets, totals, top_v, cve = _parse_excel_to_payload(path, "mitkat", meta)
            self.assertEqual(totals["high"], 1)
            self.assertEqual(len(top_v), 1)
            self.assertEqual(assets[0]["hostname"], "host1")
        finally:
            try:
                Path(path).unlink(missing_ok=True)
            except PermissionError:
                pass  # Windows may keep a brief lock after openpyxl closes the workbook.

    def test_settings_patch_adds_update_fields(self):
        from Automation.backend.exec_summary import _patch_docx_after_save

        from Automation.backend.exec_summary import _build_document

        meta = {
            "client_name": "C",
            "org_name": "O",
            "report_date": "2025-01-01",
            "assessment_type": "T",
            "document_id": "",
        }
        doc = _build_document(meta, [], {"critical": 0, "high": 0, "medium": 0, "low": 0, "informational": 0, "total": 0}, _sample_ai_payload())
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            path = tmp.name
        try:
            doc.save(path)
            _patch_docx_after_save(path)
            with zipfile.ZipFile(path, "r") as zf:
                st = zf.read("word/settings.xml").decode("utf-8")
            self.assertIn("updateFields", st)
        finally:
            Path(path).unlink(missing_ok=True)

    def test_exec_summary_routes_protected_by_session_middleware(self):
        """Production: /exec-summary requires SSO (session_security protected_paths)."""
        from unittest.mock import MagicMock

        from session_security import SessionSecurityMiddleware

        mw = SessionSecurityMiddleware(MagicMock())
        for path in ("/exec-summary", "/exec-summary/", "/exec-summary/generate/"):
            req = MagicMock()
            req.url.path = path
            self.assertTrue(
                mw._is_protected_endpoint(req),
                f"expected {path!r} to require authentication",
            )


if __name__ == "__main__":
    unittest.main()
